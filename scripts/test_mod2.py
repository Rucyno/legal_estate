# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
src_file = os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")
doc = docx.Document(src_file)

def set_run_ru(run, font_name="Times New Roman", size_pt=10, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    # Check if w:lang exists
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

# 1. Update Table 0 (date of contract)
tbl0 = doc.tables[0]
p_date = tbl0.cell(0, 1).paragraphs[0]
p_date.text = ""
r_d = p_date.add_run("«23» октября 2026 года")
set_run_ru(r_d, size_pt=9.5, bold=True)

# 2. Update P4 (Buyer preamble)
p4 = doc.paragraphs[4]
p4.text = ""
p4.paragraph_format.first_line_indent = Pt(21.25)
p4.paragraph_format.space_after = Pt(2.0)
p4.paragraph_format.line_spacing = 1.0791666666666666
r4_1 = p4.add_run("и гр. ")
set_run_ru(r4_1)
r4_2 = p4.add_run("Медведев Наиль Ядкярович")
set_run_ru(r4_2, bold=True)
r4_3 = p4.add_run(", 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина Российской Федерации: серия ")
set_run_ru(r4_3)
r4_4 = p4.add_run("5624 № 808999")
set_run_ru(r4_4, bold=True)
r4_5 = p4.add_run(", выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: ")
set_run_ru(r4_5)
r4_6 = p4.add_run("580-005")
set_run_ru(r4_6, bold=True)
r4_7 = p4.add_run(", зарегистрированный по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3, именуемый в дальнейшем «")
set_run_ru(r4_7)
r4_8 = p4.add_run("Покупатель")
set_run_ru(r4_8, bold=True)
r4_9 = p4.add_run("», с другой стороны,")
set_run_ru(r4_9)

# 3. Insert Clause 3.6 after 3.5 (before P32)
p32 = doc.paragraphs[32]
p36 = p32.insert_paragraph_before()
p36.paragraph_format.first_line_indent = Pt(21.25)
p36.paragraph_format.space_after = Pt(2.0)
p36.paragraph_format.line_spacing = 1.0791666666666666
r36 = p36.add_run("3.6. По соглашению Сторон в отчуждаемом «Объекте» передается и остается следующее имущество: в ванной комнате — тумба с раковиной, смесители; в комнате — люстра и потолочные светильники. Стоимость указанного имущества включена в общую цену «Объекта» по настоящему Договору.")
set_run_ru(r36)

# 4. Update Table 1 (DKP Signatures)
tbl1 = doc.tables[1]
# Seller date
p_s_date = tbl1.cell(0, 0).paragraphs[6] # wait, check paragraphs in cell 0
print("Cell 0 paragraphs:")
for idx, p in enumerate(tbl1.cell(0, 0).paragraphs):
    print(f"  {idx}: {repr(p.text)}")

# Buyer cell
print("Cell 1 paragraphs:")
for idx, p in enumerate(tbl1.cell(0, 1).paragraphs):
    print(f"  {idx}: {repr(p.text)}")
