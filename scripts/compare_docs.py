# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
f_orig = os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")
f_new = os.path.join(base, "2026_10_23_ДКП_с_Распиской_Лазутина_Медведев.docx")

d_orig = docx.Document(f_orig)
d_new = docx.Document(f_new)

print("Orig paragraphs count:", len(d_orig.paragraphs))
print("New paragraphs count:", len(d_new.paragraphs))
print("Orig tables count:", len(d_orig.tables))
print("New tables count:", len(d_new.tables))
