# -*- coding: utf-8 -*-
import docx, sys, glob
files = glob.glob(r"D:\*\legal_estate\deals\07_*\готовые\Основной_договор_с_распиской_*.docx")
if not files:
    # try direct search
    import os
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) if '__file__' in locals() else r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
    files = [os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")]
doc = docx.Document(files[0])
t1 = doc.tables[1]
for r_i, r in enumerate(t1.rows):
    for c_i, c in enumerate(r.cells):
        print(f"Cell ({r_i}, {c_i}) width={c.width}:")
        for p in c.paragraphs:
            print(f"   {repr(p.text)}")
