# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

def set_run_ru(run, font_name="Times New Roman", size_pt=9.0, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "2026_10_26_Договор_аренды_Греченкова_Артамонов_ОБЕСПЕЧИТЕЛЬНЫЙ.docx")
doc = docx.Document(f)

t1 = doc.tables[1]
c1 = t1.cell(0, 1)

p5 = c1.paragraphs[5]
p5.text = ""
r1 = p5.add_run("Выдан: УМВД России по Костромской области 14.08.2024 г., к/п 440-005")
set_run_ru(r1, size_pt=9.0)

# insert p6 before p_rs
p_rs = c1.paragraphs[6]
p_birth = p_rs.insert_paragraph_before()
p_birth.paragraph_format.space_after = Pt(2.0)
p_birth.paragraph_format.line_spacing = 1.05
r2 = p_birth.add_run("Дата рождения: 18.09.1982 г., гор. Кострома")
set_run_ru(r2, size_pt=9.0)

doc.save(f)
doc.save(os.path.join(base, "Договор_аренды_Греченкова_Артамонов_УСИЛЕННЫЙ_ПРОЕКТ.docx"))
print("Saved without soft break!")
