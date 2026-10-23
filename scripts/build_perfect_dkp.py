# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
src_file = os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")
doc = docx.Document(src_file)

def set_run_ru(run, font_name="Times New Roman", size_pt=10, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

# 1. Update Table 0 (date of contract)
tbl0 = doc.tables[0]
p_date = tbl0.cell(0, 1).paragraphs[0]
p_date.text = ""
r_d = p_date.add_run("«23» октября 2026 года")
set_run_ru(r_d, size_pt=9.5, bold=True)

# 2. Update P4 (Buyer preamble)
p4 = doc.paragraphs[4]
p4.text = ""
p4.paragraph_format.first_line_indent = Pt(21.25)
p4.paragraph_format.space_after = Pt(2.0)
p4.paragraph_format.line_spacing = 1.0791666666666666
r4_1 = p4.add_run("и гр. ")
set_run_ru(r4_1)
r4_2 = p4.add_run("Медведев Наиль Ядкярович")
set_run_ru(r4_2, bold=True)
r4_3 = p4.add_run(", 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина Российской Федерации: серия ")
set_run_ru(r4_3)
r4_4 = p4.add_run("5624 № 808999")
set_run_ru(r4_4, bold=True)
r4_5 = p4.add_run(", выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: ")
set_run_ru(r4_5)
r4_6 = p4.add_run("580-005")
set_run_ru(r4_6, bold=True)
r4_7 = p4.add_run(", зарегистрированный по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3, именуемый в дальнейшем «")
set_run_ru(r4_7)
r4_8 = p4.add_run("Покупатель")
set_run_ru(r4_8, bold=True)
r4_9 = p4.add_run("», с другой стороны,")
set_run_ru(r4_9)

# 3. Insert Clause 3.6 after 3.5 (before P32)
p32 = doc.paragraphs[32]
p36 = p32.insert_paragraph_before()
p36.paragraph_format.first_line_indent = Pt(21.25)
p36.paragraph_format.space_after = Pt(2.0)
p36.paragraph_format.line_spacing = 1.0791666666666666
r36 = p36.add_run("3.6. По соглашению Сторон в отчуждаемом «Объекте» передается и остается следующее имущество: в ванной комнате — тумба с раковиной, смесители; в комнате — люстра и потолочные светильники. Стоимость указанного имущества включена в общую цену «Объекта» по настоящему Договору.")
set_run_ru(r36)

# 4. Update Table 1 (DKP Signatures)
tbl1 = doc.tables[1]
# Seller date (paragraph 9)
p_s_date = tbl1.cell(0, 0).paragraphs[9]
p_s_date.text = ""
r_sd = p_s_date.add_run("Дата: «23» октября 2026 г.")
set_run_ru(r_sd, size_pt=9.5)

# Buyer cell
c1 = tbl1.cell(0, 1)
c1.paragraphs[0].text = "ПОКУПАТЕЛЬ:"
c1.paragraphs[0].runs[0].bold = True
set_run_ru(c1.paragraphs[0].runs[0], size_pt=9.5, bold=True)

c1.paragraphs[1].text = "Медведев Наиль Ядкярович"
set_run_ru(c1.paragraphs[1].runs[0], size_pt=9.5, bold=True)

c1.paragraphs[2].text = "Дата рождения: 12.04.1984 г."
set_run_ru(c1.paragraphs[2].runs[0], size_pt=9.0)

c1.paragraphs[3].text = "Паспорт РФ: 5624 № 808999"
set_run_ru(c1.paragraphs[3].runs[0], size_pt=9.0)

c1.paragraphs[4].text = "Выдан: УМВД России по Пензенской области 25.12.2024 г., к/п 580-005"
set_run_ru(c1.paragraphs[4].runs[0], size_pt=9.0)

c1.paragraphs[5].text = "Регистрация: Пензенская обл., Городищенский р-н, с. Верхняя Елюзань, ул. Заречная, д. 3"
set_run_ru(c1.paragraphs[5].runs[0], size_pt=9.0)

c1.paragraphs[7].text = "Подпись: ______________________ / Н.Я. Медведев /"
set_run_ru(c1.paragraphs[7].runs[0], size_pt=9.5)

c1.paragraphs[9].text = "Дата: «23» октября 2026 г."
set_run_ru(c1.paragraphs[9].runs[0], size_pt=9.5)

# Also ensure seller cell fonts
for p in tbl1.cell(0, 0).paragraphs:
    for r in p.runs:
        set_run_ru(r, size_pt=r.font.size.pt if r.font.size else 9.5, bold=r.bold)

# 5. Update Table 2 (Date in Raspiska header table)
tbl2 = doc.tables[2]
p_r_date = tbl2.cell(0, 1).paragraphs[0]
p_r_date.text = ""
r_rd = p_r_date.add_run("«23» октября 2026 года")
set_run_ru(r_rd, size_pt=9.5, bold=True)

# 6. Update Raspiska paragraphs (find where buyer is in raspiska)
# In raspiska:
# P50 (title date): "в получении полной оплаты стоимости квартиры по Договору купли-продажи квартиры от «___» ____________ 2026 года"
for p in doc.paragraphs:
    if "в получении полной оплаты стоимости квартиры" in p.text:
        p.text = "в получении полной оплаты стоимости квартиры по Договору купли-продажи квартиры от «23» октября 2026 года"
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_run_ru(r, size_pt=9.5, bold=True)
    if "гр. Умеркаевой Нажии Ядкяровны" in p.text:
        p.text = ""
        p.paragraph_format.first_line_indent = Pt(21.25)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.line_spacing = 1.0791666666666666
        r_rp1 = p.add_run("гр. ")
        set_run_ru(r_rp1)
        r_rp2 = p.add_run("Медведева Наиля Ядкяровича")
        set_run_ru(r_rp2, bold=True)
        r_rp3 = p.add_run(", 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина Российской Федерации: серия 5624 № 808999, выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: 580-005, зарегистрированного по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3,")
        set_run_ru(r_rp3)
    if "в соответствии с условиями Договора купли-продажи квартиры от «___»" in p.text:
        p.text = p.text.replace("«___» ____________ 2026 года", "«23» октября 2026 года")
        for r in p.runs:
            set_run_ru(r)
    if "Умеркаевой Нажие Ядкяровне" in p.text:
        p.text = p.text.replace("Умеркаевой Нажие Ядкяровне", "Медведеву Наилю Ядкяровичу")
        for r in p.runs:
            set_run_ru(r)

# 7. Update Table 3 (Signatures in Raspiska)
tbl3 = doc.tables[3]
# Seller date (paragraph 6)
p_rs_date = tbl3.cell(0, 0).paragraphs[6]
p_rs_date.text = "Дата: «23» октября 2026 г."
set_run_ru(p_rs_date.runs[0], size_pt=9.5)

# Buyer cell in raspiska
c3_1 = tbl3.cell(0, 1)
c3_1.paragraphs[0].text = "ДЕНЬГИ ПЕРЕДАЛ (ПОКУПАТЕЛЬ):"
set_run_ru(c3_1.paragraphs[0].runs[0], size_pt=9.5, bold=True)

c3_1.paragraphs[2].text = "Медведев Наиль Ядкярович"
set_run_ru(c3_1.paragraphs[2].runs[0], size_pt=9.5, bold=True)

c3_1.paragraphs[4].text = "Подпись: ______________________ / Н.Я. Медведев /"
set_run_ru(c3_1.paragraphs[4].runs[0], size_pt=9.5)

c3_1.paragraphs[6].text = "Дата: «23» октября 2026 г."
set_run_ru(c3_1.paragraphs[6].runs[0], size_pt=9.5)

for cell in tbl3.rows[0].cells:
    for p in cell.paragraphs:
        for r in p.runs:
            set_run_ru(r, size_pt=r.font.size.pt if r.font.size else 9.5, bold=r.bold)

# Save updated DKP
out_dkp = os.path.join(base, "2026_10_23_ДКП_с_Распиской_Лазутина_Медведев.docx")
doc.save(out_dkp)
print("Saved beautiful DKP with original table signatures:", out_dkp)
