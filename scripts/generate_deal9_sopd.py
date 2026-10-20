import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_sopd_document(person_data, output_path, logo_path='templates/assets/etagi_logo.png'):
    doc = docx.Document()
    
    # 1. Поля страницы: ровно на 1 лист А4
    for section in doc.sections:
        section.top_margin = Inches(0.47)     # ~1.2 см
        section.bottom_margin = Inches(0.47)  # ~1.2 см
        section.left_margin = Inches(0.59)    # ~1.5 см
        section.right_margin = Inches(0.59)   # ~1.5 см
        section.page_width = Inches(8.27)     # A4
        section.page_height = Inches(11.69)
        
    # Логотип «Этажи»
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(2)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(1.8))
        
    # Заголовок
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.line_spacing = 1.0
    r1 = p_title.add_run('СОГЛАСИЕ\n')
    r1.bold = True
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r2 = p_title.add_run('на обработку персональных данных')
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    
    # Вводная часть с реквизитами субъекта
    p_preamble = doc.add_paragraph()
    p_preamble.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_preamble.paragraph_format.first_line_indent = Inches(0.3)
    p_preamble.paragraph_format.space_before = Pt(2)
    p_preamble.paragraph_format.space_after = Pt(2)
    p_preamble.paragraph_format.line_spacing = 1.03
    
    r_pre = p_preamble.add_run(person_data['preamble_text'])
    r_pre.font.name = 'Times New Roman'
    r_pre.font.size = Pt(8.5)
    
    # 9 пунктов согласия
    clauses = [
        ('1. ', 'Данное Согласие дается на обработку персональных данных, как без использования средств автоматизации, так и с их использованием.'),
        ('2. ', 'Согласие дается на обработку перечня следующих персональных данных, не являющихся специальными или биометрическими: фамилия, имя, отчество; дата и место рождения; гражданство; реквизиты документа, удостоверяющего личность; адрес места жительства (регистрации) и пребывания; контактный номер телефона; адрес электронной почты; ИНН; СНИЛС; сведения о зарегистрированных правах на объект недвижимого имущества и сведения о переходе прав.'),
        ('3. ', 'Цель обработки персональных данных: оказание услуг агентством недвижимости, содействие в подготовке, заключении и исполнении сделок с недвижимостью, заключение с Оператором любых договоров и соглашений на оказание таких услуг, а также проверка достоверности указанных сведений.'),
        ('4. ', 'Основанием для обработки персональных данных является: ст. 24 Конституции РФ; Федеральный закон № 152-ФЗ «О персональных данных»; Политика обработки персональных данных Оператора; настоящее Согласие.'),
        ('5. ', 'В ходе обработки с персональными данными будут совершены следующие действия: сбор, запись, систематизация, накопление, хранение, уточнение (обновление, изменение), извлечение, использование, передача (предоставление, доступ) третьим лицам, обезличивание, блокирование, удаление, уничтожение.'),
        ('6. ', 'Субъект персональных данных согласен с тем, что в целях надлежащего исполнения сделок его персональные данные могут быть переданы третьим лицам (в том числе банкам, кредитным организациям, органам регистрации прав Росреестра, МФЦ, нотариусам, страховым компаниям, оценочным организациям) путем предоставления.'),
        ('7. ', 'Персональные данные обрабатываются до момента отзыва субъектом персональных данных настоящего согласия в письменной форме либо достижения целей обработки. Хранение документов осуществляется в соответствии с законодательством РФ об архивном деле.'),
        ('8. ', 'Настоящее согласие может быть отозвано субъектом персональных данных путем направления письменного заявления оператору по адресу местонахождения или на электронную почту: penzaetazhi@yandex.ru.'),
        ('9. ', 'Настоящее согласие считается предоставленным с момента подписания и действует бессрочно до момента его отзыва в установленном законом порядке.')
    ]
    
    for num, cl_text in clauses:
        p_cl = doc.add_paragraph()
        p_cl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_cl.paragraph_format.first_line_indent = Inches(0.3)
        p_cl.paragraph_format.space_before = Pt(0)
        p_cl.paragraph_format.space_after = Pt(1.5)
        p_cl.paragraph_format.line_spacing = 1.03
        
        r_num = p_cl.add_run(num)
        r_num.bold = True
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(8.5)
        
        r_body = p_cl.add_run(cl_text)
        r_body.font.name = 'Times New Roman'
        r_body.font.size = Pt(8.5)
        
    # Таблица подписи без рамок
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)
    
    cell_date = table.cell(0, 0)
    cell_sign = table.cell(0, 1)
    
    cell_date.width = Inches(2.6)
    cell_sign.width = Inches(4.5)
    
    set_cell_margins(cell_date, top=40, bottom=40, left=0, right=20)
    set_cell_margins(cell_sign, top=40, bottom=40, left=20, right=0)
    
    p_date = cell_date.paragraphs[0]
    p_date.paragraph_format.space_before = Pt(4)
    p_date.paragraph_format.space_after = Pt(0)
    r_d = p_date.add_run('Дата: «20» октября 2026 г.')
    r_d.font.name = 'Times New Roman'
    r_d.font.size = Pt(9.0)
    r_d.bold = True
    
    p_sign = cell_sign.paragraphs[0]
    p_sign.paragraph_format.space_before = Pt(4)
    p_sign.paragraph_format.space_after = Pt(0)
    r_s = p_sign.add_run('Подпись: ______________________ / ' + person_data['sign_short'] + ' /')
    r_s.font.name = 'Times New Roman'
    r_s.font.size = Pt(9.0)
    r_s.bold = True
    
    doc.save(output_path)
    print(f'Saved: {output_path}')


def create_sopd_bundle(persons, output_path, logo_path='templates/assets/etagi_logo.png'):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.47)
        section.bottom_margin = Inches(0.47)
        section.left_margin = Inches(0.59)
        section.right_margin = Inches(0.59)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    clauses = [
        ('1. ', 'Данное Согласие дается на обработку персональных данных, как без использования средств автоматизации, так и с их использованием.'),
        ('2. ', 'Согласие дается на обработку перечня следующих персональных данных, не являющихся специальными или биометрическими: фамилия, имя, отчество; дата и место рождения; гражданство; реквизиты документа, удостоверяющего личность; адрес места жительства (регистрации) и пребывания; контактный номер телефона; адрес электронной почты; ИНН; СНИЛС; сведения о зарегистрированных правах на объект недвижимого имущества и сведения о переходе прав.'),
        ('3. ', 'Цель обработки персональных данных: оказание услуг агентством недвижимости, содействие в подготовке, заключении и исполнении сделок с недвижимостью, заключение с Оператором любых договоров и соглашений на оказание таких услуг, а также проверка достоверности указанных сведений.'),
        ('4. ', 'Основанием для обработки персональных данных является: ст. 24 Конституции РФ; Федеральный закон № 152-ФЗ «О персональных данных»; Политика обработки персональных данных Оператора; настоящее Согласие.'),
        ('5. ', 'В ходе обработки с персональными данными будут совершены следующие действия: сбор, запись, систематизация, накопление, хранение, уточнение (обновление, изменение), извлечение, использование, передача (предоставление, доступ) третьим лицам, обезличивание, блокирование, удаление, уничтожение.'),
        ('6. ', 'Субъект персональных данных согласен с тем, что в целях надлежащего исполнения сделок его персональные данные могут быть переданы третьим лицам (в том числе банкам, кредитным организациям, органам регистрации прав Росреестра, МФЦ, нотариусам, страховым компаниям, оценочным организациям) путем предоставления.'),
        ('7. ', 'Персональные данные обрабатываются до момента отзыва субъектом персональных данных настоящего согласия в письменной форме либо достижения целей обработки. Хранение документов осуществляется в соответствии с законодательством РФ об архивном деле.'),
        ('8. ', 'Настоящее согласие может быть отозвано субъектом персональных данных путем направления письменного заявления оператору по адресу местонахождения или на электронную почту: penzaetazhi@yandex.ru.'),
        ('9. ', 'Настоящее согласие считается предоставленным с момента подписания и действует бессрочно до момента его отзыва в установленном законом порядке.')
    ]
    
    first = True
    for person_data in persons:
        if not first:
            p_break = doc.add_paragraph()
            p_break.paragraph_format.space_before = Pt(0)
            p_break.paragraph_format.space_after = Pt(0)
            r_br = p_break.add_run()
            r_br.add_break(docx.enum.text.WD_BREAK.PAGE)
        first = False
        
        if os.path.exists(logo_path):
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.paragraph_format.space_before = Pt(0)
            p_logo.paragraph_format.space_after = Pt(2)
            run_logo = p_logo.add_run()
            run_logo.add_picture(logo_path, width=Inches(1.8))
            
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(2)
        p_title.paragraph_format.space_after = Pt(2)
        p_title.paragraph_format.line_spacing = 1.0
        r1 = p_title.add_run('СОГЛАСИЕ\n')
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r2 = p_title.add_run('на обработку персональных данных')
        r2.bold = True
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10)
        
        p_preamble = doc.add_paragraph()
        p_preamble.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_preamble.paragraph_format.first_line_indent = Inches(0.3)
        p_preamble.paragraph_format.space_before = Pt(2)
        p_preamble.paragraph_format.space_after = Pt(2)
        p_preamble.paragraph_format.line_spacing = 1.03
        
        r_pre = p_preamble.add_run(person_data['preamble_text'])
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(8.5)
        
        for num, cl_text in clauses:
            p_cl = doc.add_paragraph()
            p_cl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_cl.paragraph_format.first_line_indent = Inches(0.3)
            p_cl.paragraph_format.space_before = Pt(0)
            p_cl.paragraph_format.space_after = Pt(1.5)
            p_cl.paragraph_format.line_spacing = 1.03
            
            r_num = p_cl.add_run(num)
            r_num.bold = True
            r_num.font.name = 'Times New Roman'
            r_num.font.size = Pt(8.5)
            
            r_body = p_cl.add_run(cl_text)
            r_body.font.name = 'Times New Roman'
            r_body.font.size = Pt(8.5)
            
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        tblPr = table._tbl.tblPr
        tblBorders = OxmlElement('w:tblBorders')
        for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            node = OxmlElement(f'w:{b}')
            node.set(qn('w:val'), 'none')
            tblBorders.append(node)
        tblPr.append(tblBorders)
        
        cell_date = table.cell(0, 0)
        cell_sign = table.cell(0, 1)
        
        cell_date.width = Inches(2.6)
        cell_sign.width = Inches(4.5)
        
        set_cell_margins(cell_date, top=40, bottom=40, left=0, right=20)
        set_cell_margins(cell_sign, top=40, bottom=40, left=20, right=0)
        
        p_date = cell_date.paragraphs[0]
        p_date.paragraph_format.space_before = Pt(4)
        p_date.paragraph_format.space_after = Pt(0)
        r_d = p_date.add_run('Дата: «20» октября 2026 г.')
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(9.0)
        r_d.bold = True
        
        p_sign = cell_sign.paragraphs[0]
        p_sign.paragraph_format.space_before = Pt(4)
        p_sign.paragraph_format.space_after = Pt(0)
        r_s = p_sign.add_run('Подпись: ______________________ / ' + person_data['sign_short'] + ' /')
        r_s.font.name = 'Times New Roman'
        r_s.font.size = Pt(9.0)
        r_s.bold = True

    doc.save(output_path)
    print(f'Saved bundle: {output_path}')


if __name__ == '__main__':
    # 1. Продавец: Ермолаева Ольга Андриановна
    preamble_ermolaeva = (
        'Я, гр. Ермолаева Ольга Андриановна, 04.12.1961 года рождения, место рождения: гор. Пенза, '
        'гражданство: Российская Федерация, паспорт гражданина РФ серия: 56 06 № 626958, выдан 09.01.2007 г. '
        'Отделом внутренних дел Ленинского района г. Пензы, код подразделения: 582-001, '
        'адрес регистрации: Пензенская область, г. Пенза, пр-д 2-й Объединенный, д. 44, кв. 1, '
        'в соответствии со ст. 9 Федерального закона от 27.07.2006 № 152-ФЗ «О персональных данных», '
        'действуя свободно, своей волей и в своем интересе, а также подтверждая свою дееспособность, '
        'даю настоящее конкретное, предметное, информированное, сознательное и однозначное Согласие '
        'на обработку персональных данных (далее – Согласие) оператору персональных данных '
        'ООО «Межрегиональный Центр Недвижимости» (далее – Оператор), местом нахождения которого является: '
        '440066, г. Пенза, пр-кт Строителей, д. 152, кв. 19, на следующих условиях:'
    )
    
    # 2. Покупатель 1: Гомонов Станислав Викторович
    preamble_gomonov = (
        'Я, гр. Гомонов Станислав Викторович, 02.06.1971 года рождения, место рождения: гор. Борзя Читинской обл., '
        'гражданство: Российская Федерация, паспорт гражданина РФ серия: 56 16 № 349795, выдан 09.06.2016 г. '
        'ОУФМС России по Пензенской области в гор. Пензе, код подразделения: 580-036, '
        'адрес регистрации: Пензенская область, г. Пенза, Военный городок, д. 137, кв. 36, '
        'в соответствии со ст. 9 Федерального закона от 27.07.2006 № 152-ФЗ «О персональных данных», '
        'действуя свободно, своей волей и в своем интересе, а также подтверждая свою дееспособность, '
        'даю настоящее конкретное, предметное, информированное, сознательное и однозначное Согласие '
        'на обработку персональных данных (далее – Согласие) оператору персональных данных '
        'ООО «Межрегиональный Центр Недвижимости» (далее – Оператор), местом нахождения которого является: '
        '440066, г. Пенза, пр-кт Строителей, д. 152, кв. 19, на следующих условиях:'
    )
    
    # 3. Покупатель 2: Андреева Галия Фяретовна
    preamble_andreeva = (
        'Я, гр. Андреева Галия Фяретовна, 12.12.1968 года рождения, место рождения: гор. Каменка Пензенской обл., '
        'гражданство: Российская Федерация, паспорт гражданина РФ серия: 56 13 № 197432, выдан 27.12.2013 г. '
        'ОУФМС России по Пензенской области в Первомайском районе гор. Пензы, код подразделения: 580-003, '
        'адрес регистрации: Пензенская область, г. Пенза, ул. Военный городок, д. 137, кв. 36, '
        'в соответствии со ст. 9 Федерального закона от 27.07.2006 № 152-ФЗ «О персональных данных», '
        'действуя свободно, своей волей и в своем интересе, а также подтверждая свою дееспособность, '
        'даю настоящее конкретное, предметное, информированное, сознательное и однозначное Согласие '
        'на обработку персональных данных (далее – Согласие) оператору персональных данных '
        'ООО «Межрегиональный Центр Недвижимости» (далее – Оператор), местом нахождения которого является: '
        '440066, г. Пенза, пр-кт Строителей, д. 152, кв. 19, на следующих условиях:'
    )

    p_ermolaeva = {
        'preamble_text': preamble_ermolaeva,
        'sign_short': 'Ермолаева О.А.'
    }
    
    p_gomonov = {
        'preamble_text': preamble_gomonov,
        'sign_short': 'Гомонов С.В.'
    }
    
    p_andreeva = {
        'preamble_text': preamble_andreeva,
        'sign_short': 'Андреева Г.Ф.'
    }

    out_dir = 'deals/09_Ермолаева_Гомонов_Андреева/готовые'
    os.makedirs(out_dir, exist_ok=True)
    
    f_ermolaeva = os.path.join(out_dir, '2026_10_20_СОПД_Продавец_Ермолаева_Ольга_Андриановна.docx')
    f_gomonov = os.path.join(out_dir, '2026_10_20_СОПД_Покупатель_Гомонов_Станислав_Викторович.docx')
    f_andreeva = os.path.join(out_dir, '2026_10_20_СОПД_Покупатель_Андреева_Галия_Фяретовна.docx')
    f_bundle = os.path.join(out_dir, '2026_10_20_СОПД_Комплект_Ермолаева_Гомонов_Андреева_ВСЕ_СТОРОНЫ.docx')
    
    create_sopd_document(p_ermolaeva, f_ermolaeva)
    create_sopd_document(p_gomonov, f_gomonov)
    create_sopd_document(p_andreeva, f_andreeva)
    
    # Пакетный файл на 3 страницы (по 1 листу на каждого)
    create_sopd_bundle([p_ermolaeva, p_gomonov, p_andreeva], f_bundle)
    print('DEAL 9 ALL SOPD FILES GENERATED SUCCESSFULLY!')
