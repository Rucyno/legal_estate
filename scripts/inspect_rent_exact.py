# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "Договор_аренды_Греченкова_Артамонов_УСИЛЕННЫЙ_ПРОЕКТ.docx")
doc = docx.Document(f)

print("P8:", repr(doc.paragraphs[8].text))
print("P9:", repr(doc.paragraphs[9].text))
print("P17:", repr(doc.paragraphs[17].text))
print("P18:", repr(doc.paragraphs[18].text))
print("P19:", repr(doc.paragraphs[19].text))
print("P22:", repr(doc.paragraphs[22].text))
print("P23:", repr(doc.paragraphs[23].text))
