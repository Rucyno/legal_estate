# -*- coding: utf-8 -*-
import docx, os

folder = r"C:\Users\PC\Desktop\Тёма"
files = [f for f in os.listdir(folder) if f.endswith(".docx") and not f.startswith("~$")]
print("Files found:", files)

with open("scan_all_dates.txt", "w", encoding="utf-8") as out:
    for fname in files:
        fpath = os.path.join(folder, fname)
        doc = docx.Document(fpath)
        out.write(f"\n=================== {fname} ===================\n")
        for i, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            if any(w in txt.lower() for w in ["октябр", "13.", "14.", "2026"]):
                out.write(f"P{i}: {txt}\n")
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    txt = cell.text.strip()
                    if any(w in txt.lower() for w in ["октябр", "13.", "14.", "2026"]):
                        out.write(f"Table{t_idx} [{r_idx},{c_idx}]: {txt}\n")

print("Scan complete. Output written to scan_all_dates.txt")
