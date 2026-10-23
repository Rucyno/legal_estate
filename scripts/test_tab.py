# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

# Let us test tab stop signature
doc = docx.Document()
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
tabs = parse_xml(f"""
<w:tabs {nsdecls("w")}>
    <w:tab w:val="right" w:pos="9640"/>
</w:tabs>
""")
pPr.append(tabs)
r1 = p.add_run("ФИО")
r1.font.italic = True
r1.font.size = Pt(8)
r2 = p.add_run("\tподпись")
r2.font.italic = True
r2.font.size = Pt(8)
doc.save("test_tab.docx")
print("Saved test_tab.docx")
