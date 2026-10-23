# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

def set_run_ru(run, font_name="Times New Roman", size_pt=10, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "2026_10_26_Договор_аренды_Греченкова_Артамонов_ОБЕСПЕЧИТЕЛЬНЫЙ.docx")
doc = docx.Document(f)

# 1. Update P3 (preamble)
p3 = doc.paragraphs[3]
p3.text = ""
p3.paragraph_format.first_line_indent = Pt(21.25)
p3.paragraph_format.space_after = Pt(2.0)
p3.paragraph_format.line_spacing = 1.0791666666666666

r3_1 = p3.add_run("Индивидуальный предприниматель ")
set_run_ru(r3_1)
r3_2 = p3.add_run("Артамонов Алексей Александрович")
set_run_ru(r3_2, bold=True)
r3_3 = p3.add_run(", 18.09.1982 года рождения, пол: мужской, место рождения: гор. Кострома, ОГРНИП: ")
set_run_ru(r3_3)
r3_4 = p3.add_run("307440136000110")
set_run_ru(r3_4, bold=True)
r3_5 = p3.add_run(", ИНН: ")
set_run_ru(r3_5)
r3_6 = p3.add_run("440104330200")
set_run_ru(r3_6, bold=True)
r3_7 = p3.add_run(", действующий на основании записи в ЕГРИП, паспорт гражданина РФ: серия ")
set_run_ru(r3_7)
r3_8 = p3.add_run("34 24 № 437648")
set_run_ru(r3_8, bold=True)
r3_9 = p3.add_run(", выдан: УМВД России по Костромской области 14.08.2024 г., код подразделения: ")
set_run_ru(r3_9)
r3_10 = p3.add_run("440-005")
set_run_ru(r3_10, bold=True)
r3_11 = p3.add_run(", зарегистрированный по адресу: 216145, Смоленская область, муниципальный округ Монастырщинский, д. Колосовка, д. 15, телефон: 89159003435, e-mail: sveta4480@bk.ru, именуемый в дальнейшем «")
set_run_ru(r3_11)
r3_12 = p3.add_run("Арендатор")
set_run_ru(r3_12, bold=True)
r3_13 = p3.add_run("», с другой стороны,")
set_run_ru(r3_13)

# 2. Update Table 1 (Section 10 details)
t1 = doc.tables[1]
c1 = t1.cell(0, 1)

# Paragraph 4: Passport series and number
p_pass = c1.paragraphs[4]
p_pass.text = ""
r_p1 = p_pass.add_run("Паспорт РФ: серия 34 24 № 437648")
set_run_ru(r_p1, size_pt=9.0)

# Paragraph 5: Issued by, date, unit code, birth date
p_iss = c1.paragraphs[5]
p_iss.text = ""
r_p2 = p_iss.add_run("Выдан: УМВД России по Костромской области 14.08.2024 г., к/п 440-005\nДата рождения: 18.09.1982 г., гор. Кострома")
set_run_ru(r_p2, size_pt=9.0)

doc.save(f)
# Also sync project docx
doc.save(os.path.join(base, "Договор_аренды_Греченкова_Артамонов_УСИЛЕННЫЙ_ПРОЕКТ.docx"))
print("Successfully filled in Artamonov passport details!")
