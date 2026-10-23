# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "2026_10_26_Договор_аренды_Греченкова_Артамонов_ОБЕСПЕЧИТЕЛЬНЫЙ.docx")
doc = docx.Document(f)

print("P3 text:\n", doc.paragraphs[3].text)
print("\nT1 C1 paragraphs:")
for idx, p in enumerate(doc.tables[1].cell(0, 1).paragraphs):
    print(f"  {idx}: {repr(p.text)}")
