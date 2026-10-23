# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
out_act = os.path.join(base, "2026_10_23_Передаточный_акт_Лазутина_Медведев.docx")

doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Pt(30.0)
    s.bottom_margin = Pt(30.0)
    s.left_margin = Pt(40.0)
    s.right_margin = Pt(40.0)

def set_run_ru(run, font_name="Times New Roman", size_pt=9.5, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

def add_p(text, size_pt=9.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=1.5, line_spacing=1.05, indent=21.25, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Pt(indent)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_ru(run, size_pt=size_pt, bold=bold, italic=italic)
    return p

# Header Title
add_p("ПЕРЕДАТОЧНЫЙ АКТ", size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=1.0)
add_p("(АКТ ПРИЕМА-ПЕРЕДАЧИ КВАРТИРЫ)", size_pt=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=4.0)

# Header Table (City, Date)
tbl_h = doc.add_table(rows=1, cols=2)
tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_h.autofit = False
for cell in tbl_h.rows[0].cells:
    cell.width = Pt(255)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r"""
        <w:tcBorders %s >
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    """ % nsdecls("w"))
    tcPr.append(tcBorders)

p_l = tbl_h.cell(0, 0).paragraphs[0]
p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_l.paragraph_format.space_after = Pt(4.0)
r_l = p_l.add_run("Российская Федерация, город Пенза")
set_run_ru(r_l, size_pt=9.5, bold=True)

p_r = tbl_h.cell(0, 1).paragraphs[0]
p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_r.paragraph_format.space_after = Pt(4.0)
r_r = p_r.add_run("«23» октября 2026 года")
set_run_ru(r_r, size_pt=9.5, bold=True)

add_p("Мы, нижеподписавшиеся:", size_pt=9.5, bold=False, indent=21.25, space_after=1.5)
add_p("гр. Лазутина Елена Викторовна, 10.09.1984 года рождения, пол: женский, место рождения: гор. Пенза, паспорт гражданина РФ: серия 56 05 № 562191, выдан: Отделом внутренних дел Первомайского района г. Пензы 05.05.2006 г., код подразделения: 582-003, зарегистрированная по адресу: Российская Федерация, Пензенская область, Пензенский район, с. Засечное, ул. Олимпийская, д. 3, кв. 129, именуемая в дальнейшем «Продавец», с одной стороны,", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("и гр. Медведев Наиль Ядкярович, 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина РФ: серия 5624 № 808999, выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: 580-005, зарегистрированный по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3, именуемый в дальнейшем «Покупатель», с другой стороны,", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("совместно именуемые «Стороны», составили настоящий Передаточный акт к Договору купли-продажи квартиры от «23» октября 2026 года о нижеследующем:", size_pt=9.5, indent=21.25, space_after=2.0)

add_p("1. В соответствии со статьей 556 Гражданского кодекса РФ и условиями Договора купли-продажи квартиры от «23» октября 2026 года Продавец передал, а Покупатель принял в собственность следующее недвижимое имущество:", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("- жилое помещение — квартиру, расположенную по адресу: Российская Федерация, Пензенская область, Пензенский район, с. Засечное, ул. Олимпийская, д. 3, кв. 129, общей площадью: 43.6 (Сорок три целых шесть десятых) кв. м, расположенную на 5 этаже жилого дома, кадастровый номер: 58:24:0381302:17841 (далее — «Объект»).", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("2. Квартира осмотрена Покупателем до подписания настоящего акта. Претензий по санитарно-техническому состоянию, отделке, дверным и оконным блокам, радиаторам отопления и инженерным коммуникациям Покупатель к Продавцу не имеет. Квартира передана в качественном состоянии, пригодном для постоянного проживания.", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("3. В соответствии с условиями Договора купли-продажи квартиры от «23» октября 2026 года в квартире передается и остается следующее имущество: в ванной комнате — тумба с раковиной, смесители; в комнате — люстра и потолочные светильники. Претензий по комплектности и состоянию указанного имущества у Покупателя не имеется.", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("4. Одновременно с передачей квартиры Продавец передал Покупателю ключи от квартиры, подъезда и домофона, а также документы об отсутствии задолженностей по ЖКУ, электроэнергии, капремонту и ТКО.", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("5. Показания приборов учета: электроэнергия № ___________, показ. _______ кВт*ч; ХВС № __________, показ. _______ куб. м; ГВС № __________, показ. _______ куб. м; газ № __________, показ. _______ куб. м.", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("6. Взаимные финансовые расчеты между Сторонами произведены в полном объеме наличными деньгами до подписания настоящего акта в сумме 4 600 000 (Четыре миллиона шестьсот тысяч) рублей 00 копеек. Претензий по оплате у Продавца нет.", size_pt=9.5, indent=21.25, space_after=1.5)
add_p("7. Обязательства Сторон по передаче, принятию и оплате квартиры исполнены полностью и надлежащим образом. Взаимных претензий Стороны друг к другу не имеют. Настоящий акт составлен в 2 подлинных экземплярах.", size_pt=9.5, indent=21.25, space_after=3.0, keep_next=True)

def add_sig_block(title_text, keep_all=True):
    # Paragraph 1: Title
    add_p(title_text, size_pt=9.5, bold=True, indent=0, space_after=1.0, keep_next=True)
    
    # Paragraph 2: Full width line (paragraph with bottom border)
    p_line = doc.add_paragraph()
    pPr = p_line._p.get_or_add_pPr()
    pBdr = parse_xml(f"""
    <w:pBdr {nsdecls("w")}>
        <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>
    </w:pBdr>
    """)
    pPr.append(pBdr)
    p_line.paragraph_format.space_after = Pt(1.5)
    p_line.paragraph_format.keep_with_next = True
    r_empty = p_line.add_run(" ")
    set_run_ru(r_empty, size_pt=9.5)
    
    # Paragraph 3: Subtext with right tab
    p_sub = doc.add_paragraph()
    pPr_sub = p_sub._p.get_or_add_pPr()
    tabs = parse_xml(f"""
    <w:tabs {nsdecls("w")}>
        <w:tab w:val="right" w:pos="9640"/>
    </w:tabs>
    """)
    pPr_sub.append(tabs)
    p_sub.paragraph_format.space_after = Pt(6.0 if keep_all else 2.0)
    if keep_all:
        p_sub.paragraph_format.keep_with_next = True
    
    r1 = p_sub.add_run("ФИО")
    set_run_ru(r1, size_pt=8.0, italic=True)
    r2 = p_sub.add_run("\tподпись")
    set_run_ru(r2, size_pt=8.0, italic=True)

add_sig_block("Сдал (Продавец):", keep_all=True)
add_sig_block("Принял (Покупатель):", keep_all=False)

doc.save(out_act)
print("Act with FULL-WIDTH BORDER SIGNATURES saved successfully:", out_act)
