# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "Договор_аренды_Греченкова_Артамонов_УСИЛЕННЫЙ_ПРОЕКТ.docx")
doc = docx.Document(f)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if any(k in txt for k in ["гарантийн", "обеспечительн", "19 000", "38 000", "оплат", "плат", "2026", "срок", "Акт"]):
        print(f"P{i}: {txt[:100]}")
