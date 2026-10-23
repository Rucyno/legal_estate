# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "2026_10_26_Договор_аренды_Греченкова_Артамонов_ОБЕСПЕЧИТЕЛЬНЫЙ.docx")
doc = docx.Document(f)

for i in range(len(doc.paragraphs)):
    txt = doc.paragraphs[i].text
    if "Артамонов" in txt or "паспорт" in txt.lower() or "серия" in txt.lower():
        print(f"P{i}: {txt}\n")

t1 = doc.tables[1]
print("T1 Cell 1 paragraphs:")
for idx, p in enumerate(t1.cell(0, 1).paragraphs):
    print(f"  {idx}: {repr(p.text)}")
