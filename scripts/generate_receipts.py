# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_single_receipt(is_full, out_path):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.5)
    
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2, line_spacing=1.05, indent=Cm(0.8), bold=False):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = line_spacing
        pf.first_line_indent = indent
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
        return p

    def add_top_table():
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.columns[0].width = Cm(8.75)
        t.columns[1].width = Cm(8.75)
        
        tblPr = t._tbl.tblPr
        tblBorders = parse_xml(r'''
            <w:tblBorders %s>
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''' % nsdecls('w'))
        tblPr.append(tblBorders)
        
        c0 = t.cell(0, 0)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("г. Пенза")
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.bold = True
        
        c1 = t.cell(0, 1)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run("«___» ____________ 2026 г.")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.bold = True

    if is_full:
        p_t1 = add_p("РАСПИСКА", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1, indent=0, bold=True)
        p_t1.runs[0].font.size = Pt(12)
        p_s1 = add_p("в получении денежных средств по договору купли-продажи квартиры", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, indent=0, bold=True)
        p_s1.runs[0].font.size = Pt(10.5)

        add_top_table()
        add_p("", space_before=1, space_after=2, indent=0)

        p1 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p1.add_run("Я, гр. ").font.name = 'Times New Roman'
        p1.add_run("НУРЫЕВ ЧАРЫЯР").bold = True
        p1.add_run(", 18.12.1971 года рождения, место рождения: ТУРКМЕНИСТАН, гражданство: ТУРКМЕНИСТАН, пол: мужской, паспорт: А2859533, выдан ГОСУДАРСТВЕННАЯ МИГРАЦИОННАЯ СЛУЖБА ТУРКМЕНИСТАНА 18.09.2024 года, вид на жительство иностранного гражданина 83 № 0618724, выдан УМВД РОССИИ ПО ПЕНЗЕНСКОЙ ОБЛАСТИ 14.03.2023 года, зарегистрированный по адресу: Пензенская обл., г. Пенза, ул. Мира, д. 67, кв. 340, именуемый в дальнейшем «Продавец»,")

        p2 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2.add_run("настоящей распиской подтверждаю, что получил от гр. ").font.name = 'Times New Roman'
        p2.add_run("ЛУКАШОВА ЕГОРА ВЯЧЕСЛАВОВИЧА").bold = True
        p2.add_run(", 01.07.2003 года рождения, место рождения: Г. ПЕНЗА ПЕНЗЕНСКАЯ ОБЛ., гражданство: Российская Федерация, пол: мужской, паспорт: серия 5623 № 732738, код подразделения: 580-005, выдан УМВД РОССИИ ПО ПЕНЗЕНСКОЙ ОБЛАСТИ 22.08.2023 года, зарегистрированного по адресу: обл. Пензенская, р-н Городищенский, с. Канаевка, ул. Молодежная, двлд. 95, именуемого в дальнейшем «Покупатель»,")

        p3 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p3.add_run("денежные средства в размере ").font.name = 'Times New Roman'
        p3.add_run("11 300 000 (Одиннадцать миллионов триста тысяч) рублей 00 копеек").bold = True
        p3.add_run(" в качестве полной оплаты стоимости следующего недвижимого имущества: ")
        p3.add_run("квартира, кадастровый номер: 58:29:4002011:1832, общей площадью 97.1 кв. м., находящаяся по адресу: Российская Федерация, Пензенская область, г.о. город Пенза, г. Пенза, ул. Мира, д. 67, кв. 340").bold = True
        p3.add_run(", проданной по Договору купли-продажи квартиры с использованием кредитных средств от «08» октября 2026 года.")

        p4 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p4.add_run("Оплата цены Квартиры произведена Покупателем в полном объеме в соответствии с разделом 2 Договора купли-продажи в следующем порядке:")

        p4_1 = add_p(space_before=0, space_after=1, line_spacing=1.05, indent=Cm(1.2))
        p4_1.add_run("1) ").font.name = 'Times New Roman'
        p4_1.add_run("100 000 (Сто тысяч) рублей 00 копеек").bold = True
        p4_1.add_run(" — оплачена Покупателем Продавцу до подписания Договора в качестве задатка (ст. 380 ГК РФ);")

        p4_2 = add_p(space_before=0, space_after=1, line_spacing=1.05, indent=Cm(1.2))
        p4_2.add_run("2) ").font.name = 'Times New Roman'
        p4_2.add_run("8 930 000 (Восемь миллионов девятьсот тридцать тысяч) рублей 00 копеек").bold = True
        p4_2.add_run(" — выплачена за счет собственных средств Покупателя посредством аккредитива, открытого в ТКБ БАНК ПАО;")

        p4_3 = add_p(space_before=0, space_after=2, line_spacing=1.05, indent=Cm(1.2))
        p4_3.add_run("3) ").font.name = 'Times New Roman'
        p4_3.add_run("2 270 000 (Два миллиона двести семьдесят тысяч) рублей 00 копеек").bold = True
        p4_3.add_run(" — выплачена за счет кредитных средств, предоставленных Публичным акционерным обществом «ТРАНСКАПИТАЛБАНК» (ТКБ БАНК ПАО) по Кредитному договору № ИК000026/01061 от «08» октября 2026 г., посредством аккредитива, открытого в ТКБ БАНК ПАО.")

        p5 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p5.add_run("Денежные средства по Договору получены мной сполна. Обязательства Покупателя по оплате стоимости Квартиры исполнены полностью и в установленные сроки. Материальных и иных финансовых претензий по оплате Квартиры к Покупателю не имею.")

        p6 = add_p(space_before=1, space_after=5, line_spacing=1.05, indent=Cm(0.8))
        p6.add_run("Настоящая расписка составлена в двух равнозначных экземплярах — по одному для каждой из Сторон.")

        p_sig1 = add_p(space_before=4, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_sig1.add_run("Продавец денежные средства получил сполна:").bold = True

        p_line1 = add_p("__________________________________________________________________________________________________", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_sub1 = add_p("(фамилия, имя, отчество полностью от руки и личная подпись Продавца)", space_before=0, space_after=2, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_sub1.runs[0].font.size = Pt(8)
        p_sub1.runs[0].italic = True

        p_name1 = add_p("Нурыев Чарыяр                                                 / ____________________ /", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_date1 = add_p("Дата: «___» ____________ 2026 г.", space_before=1, space_after=6, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_date1.runs[0].bold = True

        p_sig2 = add_p(space_before=4, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_sig2.add_run("Покупатель денежные средства передал:").bold = True

        p_line2 = add_p("__________________________________________________________________________________________________", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_sub2 = add_p("(фамилия, имя, отчество полностью от руки и личная подпись Покупателя)", space_before=0, space_after=2, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_sub2.runs[0].font.size = Pt(8)
        p_sub2.runs[0].italic = True

        p_name2 = add_p("Лукашов Егор Вячеславович                            / ____________________ /", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_date2 = add_p("Дата: «___» ____________ 2026 г.", space_before=1, space_after=0, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_date2.runs[0].bold = True
    else:
        p_t2 = add_p("РАСПИСКА", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1, indent=0, bold=True)
        p_t2.runs[0].font.size = Pt(12)
        p_s2 = add_p("в получении денежных средств (задатка) за квартиру", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6, indent=0, bold=True)
        p_s2.runs[0].font.size = Pt(10.5)

        add_top_table()
        add_p("", space_before=1, space_after=2, indent=0)

        p2_1 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2_1.add_run("Я, гр. ").font.name = 'Times New Roman'
        p2_1.add_run("НУРЫЕВ ЧАРЫЯР").bold = True
        p2_1.add_run(", 18.12.1971 года рождения, место рождения: ТУРКМЕНИСТАН, гражданство: ТУРКМЕНИСТАН, пол: мужской, паспорт: А2859533, выдан ГОСУДАРСТВЕННАЯ МИГРАЦИОННАЯ СЛУЖБА ТУРКМЕНИСТАНА 18.09.2024 года, вид на жительство иностранного гражданина 83 № 0618724, выдан УМВД РОССИИ ПО ПЕНЗЕНСКОЙ ОБЛАСТИ 14.03.2023 года, зарегистрированный по адресу: Пензенская обл., г. Пенза, ул. Мира, д. 67, кв. 340, именуемый в дальнейшем «Продавец»,")

        p2_2 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2_2.add_run("настоящей распиской подтверждаю, что получил от гр. ").font.name = 'Times New Roman'
        p2_2.add_run("ЛУКАШОВА ЕГОРА ВЯЧЕСЛАВОВИЧА").bold = True
        p2_2.add_run(", 01.07.2003 года рождения, место рождения: Г. ПЕНЗА ПЕНЗЕНСКАЯ ОБЛ., гражданство: Российская Федерация, пол: мужской, паспорт: серия 5623 № 732738, код подразделения: 580-005, выдан УМВД РОССИИ ПО ПЕНЗЕНСКОЙ ОБЛАСТИ 22.08.2023 года, зарегистрированного по адресу: обл. Пензенская, р-н Городищенский, с. Канаевка, ул. Молодежная, двлд. 95, именуемого в дальнейшем «Покупатель»,")

        p2_3 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2_3.add_run("денежные средства в размере ").font.name = 'Times New Roman'
        p2_3.add_run("100 000 (Сто тысяч) рублей 00 копеек").bold = True
        p2_3.add_run(" в качестве задатка в соответствии со ст. 380 ГК РФ и п. 2.2.1 Договора купли-продажи квартиры с использованием кредитных средств от «08» октября 2026 года в счет оплаты стоимости следующего приобретаемого недвижимого имущества: ")
        p2_3.add_run("квартира, кадастровый номер: 58:29:4002011:1832, общей площадью 97.1 кв. м., находящаяся по адресу: Российская Федерация, Пензенская область, городской округ город Пенза, г. Пенза, ул. Мира, д. 67, кв. 340").bold = True
        p2_3.add_run(".")

        p2_4 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2_4.add_run("Указанная сумма задатка в размере 100 000 (Сто тысяч) рублей 00 копеек получена мной в полном объеме до подписания Договора купли-продажи. Оставшаяся часть стоимости Квартиры в размере 11 200 000 (Одиннадцать миллионов двести тысяч) рублей 00 копеек подлежит оплате Покупателем посредством аккредитива в соответствии с условиями Договора купли-продажи.")

        p2_5 = add_p(space_before=1, space_after=2, line_spacing=1.05, indent=Cm(0.8))
        p2_5.add_run("Претензий по передаче суммы задатка к Покупателю не имею.")

        p2_6 = add_p(space_before=1, space_after=8, line_spacing=1.05, indent=Cm(0.8))
        p2_6.add_run("Настоящая расписка составлена в двух равнозначных экземплярах — по одному для каждой из Сторон.")

        p2_sig1 = add_p(space_before=6, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_sig1.add_run("Продавец денежные средства получил:").bold = True

        p2_line1 = add_p("__________________________________________________________________________________________________", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_sub1 = add_p("(фамилия, имя, отчество полностью от руки и личная подпись Продавца)", space_before=0, space_after=2, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p2_sub1.runs[0].font.size = Pt(8)
        p2_sub1.runs[0].italic = True

        p2_name1 = add_p("Нурыев Чарыяр                                                 / ____________________ /", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_date1 = add_p("Дата: «___» ____________ 2026 г.", space_before=1, space_after=8, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_date1.runs[0].bold = True

        p2_sig2 = add_p(space_before=6, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_sig2.add_run("Покупатель денежные средства передал:").bold = True

        p2_line2 = add_p("__________________________________________________________________________________________________", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_sub2 = add_p("(фамилия, имя, отчество полностью от руки и личная подпись Покупателя)", space_before=0, space_after=2, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p2_sub2.runs[0].font.size = Pt(8)
        p2_sub2.runs[0].italic = True

        p2_name2 = add_p("Лукашов Егор Вячеславович                            / ____________________ /", space_before=1, space_after=1, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_date2 = add_p("Дата: «___» ____________ 2026 г.", space_before=1, space_after=0, line_spacing=1.0, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        p2_date2.runs[0].bold = True

    doc.save(out_path)
    print("Saved:", out_path)

if __name__ == '__main__':
    create_single_receipt(True, r"C:\Users\PC\Desktop\Тёма\Расписка_Нурыев_Лукашов_11_300_000.docx")
    create_single_receipt(False, r"C:\Users\PC\Desktop\Тёма\Расписка_Нурыев_Лукашов_100_000_задаток.docx")
