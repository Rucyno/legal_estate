# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

def set_run_ru(run, font_name="Times New Roman", size_pt=9.5, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
src_file = os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")
doc = docx.Document(src_file)

# 1. Update Table 0 (date of contract)
tbl0 = doc.tables[0]
p_date = tbl0.cell(0, 1).paragraphs[0]
p_date.text = ""
r_d = p_date.add_run("«23» октября 2026 года")
set_run_ru(r_d, size_pt=9.5, bold=True)

# 2. Update P4 (Buyer preamble)
p4 = doc.paragraphs[4]
p4.text = ""
p4.paragraph_format.first_line_indent = Pt(21.25)
p4.paragraph_format.space_after = Pt(2.0)
p4.paragraph_format.line_spacing = 1.0791666666666666
r4_1 = p4.add_run("и гр. ")
set_run_ru(r4_1)
r4_2 = p4.add_run("Медведев Наиль Ядкярович")
set_run_ru(r4_2, bold=True)
r4_3 = p4.add_run(", 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина Российской Федерации: серия ")
set_run_ru(r4_3)
r4_4 = p4.add_run("5624 № 808999")
set_run_ru(r4_4, bold=True)
r4_5 = p4.add_run(", выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: ")
set_run_ru(r4_5)
r4_6 = p4.add_run("580-005")
set_run_ru(r4_6, bold=True)
r4_7 = p4.add_run(", зарегистрированный по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3, именуемый в дальнейшем «")
set_run_ru(r4_7)
r4_8 = p4.add_run("Покупатель")
set_run_ru(r4_8, bold=True)
r4_9 = p4.add_run("», с другой стороны,")
set_run_ru(r4_9)

# 2.1 Update Clause 3.2 (14 calendar days)
for p in doc.paragraphs:
    if "3.2. Продавец обязуется фактически освободить" in p.text:
        p.text = ""
        p.paragraph_format.first_line_indent = Pt(21.25)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.line_spacing = 1.0791666666666666
        r32 = p.add_run("3.2. Продавец обязуется фактически освободить «Объект» от находящегося в нем личного имущества и передать его Покупателю по Передаточному акту вместе с комплектом ключей в течение 14 (четырнадцати) календарных дней после полной оплаты Покупателем стоимости «Объекта», указанной в п. 2.1 настоящего Договора.")
        set_run_ru(r32)
        break

# 3. Insert Clause 3.6 after 3.5 (before P32)
p32 = doc.paragraphs[32]
p36 = p32.insert_paragraph_before()
p36.paragraph_format.first_line_indent = Pt(21.25)
p36.paragraph_format.space_after = Pt(2.0)
p36.paragraph_format.line_spacing = 1.0791666666666666
r36 = p36.add_run("3.6. По соглашению Сторон в отчуждаемом «Объекте» передается и остается следующее имущество: в ванной комнате — тумба с раковиной, смесители; в комнате — люстра и потолочные светильники. Стоимость указанного имущества включена в общую цену «Объекта» по настоящему Договору.")
set_run_ru(r36)

# 3.1 Update Clause 5.2 (Registered persons: Lazutina E.V. and minor Lazutin Makar Viktorovich, 10 days deregistration)
for p in doc.paragraphs:
    if "5.2. Продавец гарантирует" in p.text:
        p.text = ""
        p.paragraph_format.first_line_indent = Pt(21.25)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.line_spacing = 1.0791666666666666
        r52 = p.add_run("5.2. На момент подписания настоящего Договора в «Объекте» зарегистрированы по месту жительства: Лазутина Елена Викторовна и несовершеннолетний Лазутин Макар Викторович, 12.10.2016 года рождения. Продавец обязуется обеспечить снятие себя и несовершеннолетнего Лазутина Макара Викторовича с регистрационного учета по месту жительства из отчуждаемого «Объекта» в течение 10 (десяти) календарных дней после полной оплаты стоимости «Объекта» по настоящему Договору. Иных лиц, сохраняющих право пользования жилым помещением в соответствии со статьями 292, 558 Гражданского кодекса РФ, не имеется.")
        set_run_ru(r52)
        break

# =========================================================================
# 4. REPLACE TABLE 1 (DKP SIGNATURES) WITH FULL-WIDTH ROW FORMAT
# =========================================================================
tbl1 = doc.tables[1]
tbl1_elem = tbl1._tbl
parent1 = tbl1_elem.getparent()
idx1 = parent1.index(tbl1_elem)

def make_p_xml(text, size_pt=9.5, bold=False, italic=False, space_before=0, space_after=2, keep_next=True):
    sz_val = int(size_pt * 2)
    b_tag = "<w:b/>" if bold else ""
    i_tag = "<w:i/>" if italic else ""
    kn_tag = "<w:keepNext/>" if keep_next else ""
    sb_attr = f'w:before="{int(space_before * 20)}"' if space_before else ""
    sa_attr = f'w:after="{int(space_after * 20)}"' if space_after else ""
    sp_tag = f'<w:spacing {sb_attr} {sa_attr}/>' if (sb_attr or sa_attr) else ""
    
    # We escape XML entities in text
    safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    
    xml_str = f"""
    <w:p {nsdecls("w")}>
        <w:pPr>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                {b_tag}
                {i_tag}
                <w:sz w:val="{sz_val}"/>
                <w:lang w:val="ru-RU"/>
            </w:rPr>
            {kn_tag}
            {sp_tag}
        </w:pPr>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                {b_tag}
                {i_tag}
                <w:sz w:val="{sz_val}"/>
                <w:lang w:val="ru-RU"/>
            </w:rPr>
            <w:t xml:space="preserve">{safe_text}</w:t>
        </w:r>
    </w:p>
    """
    return parse_xml(xml_str)

line_str = "_" * 95

elems_to_insert = [
    make_p_xml("Продавец:", size_pt=9.5, bold=True, space_before=2, space_after=1, keep_next=True),
    make_p_xml(line_str, size_pt=9.5, space_before=0, space_after=0, keep_next=True),
    make_p_xml("ФИО                                                                                                                                           подпись", size_pt=8.0, italic=True, space_before=0, space_after=8, keep_next=True),
    make_p_xml("Покупатель:", size_pt=9.5, bold=True, space_before=2, space_after=1, keep_next=True),
    make_p_xml(line_str, size_pt=9.5, space_before=0, space_after=0, keep_next=True),
    make_p_xml("ФИО                                                                                                                                           подпись", size_pt=8.0, italic=True, space_before=0, space_after=4, keep_next=False),
]

for el in elems_to_insert:
    parent1.insert(idx1, el)
    idx1 += 1

parent1.remove(tbl1_elem)

# 5. Update Table 2 (Date in Raspiska header table)
tbl2 = doc.tables[1] # because we removed tbl1, tbl2 is now doc.tables[1]
p_r_date = tbl2.cell(0, 1).paragraphs[0]
p_r_date.text = ""
r_rd = p_r_date.add_run("«23» октября 2026 года")
set_run_ru(r_rd, size_pt=9.5, bold=True)

# 6. Update Raspiska paragraphs
for p in doc.paragraphs:
    if "в получении полной оплаты стоимости квартиры" in p.text:
        p.text = "в получении полной оплаты стоимости квартиры по Договору купли-продажи квартиры от «23» октября 2026 года"
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            set_run_ru(r, size_pt=9.5, bold=True)
    if "гр. Умеркаевой Нажии Ядкяровны" in p.text:
        p.text = ""
        p.paragraph_format.first_line_indent = Pt(21.25)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.line_spacing = 1.0791666666666666
        r_rp1 = p.add_run("гр. ")
        set_run_ru(r_rp1)
        r_rp2 = p.add_run("Медведева Наиля Ядкяровича")
        set_run_ru(r_rp2, bold=True)
        r_rp3 = p.add_run(", 12.04.1984 года рождения, пол: мужской, место рождения: село Верхняя Елюзань Городищенского района Пензенской области, паспорт гражданина Российской Федерации: серия 5624 № 808999, выдан: УМВД России по Пензенской области 25.12.2024 г., код подразделения: 580-005, зарегистрированного по адресу: Российская Федерация, Пензенская область, Городищенский район, с. Верхняя Елюзань, ул. Заречная, д. 3,")
        set_run_ru(r_rp3)
    if "в соответствии с условиями Договора купли-продажи квартиры от «___»" in p.text:
        p.text = p.text.replace("«___» ____________ 2026 года", "«23» октября 2026 года")
        for r in p.runs:
            set_run_ru(r)
    if "Умеркаевой Нажие Ядкяровне" in p.text:
        p.text = p.text.replace("Умеркаевой Нажие Ядкяровне", "Медведеву Наилю Ядкяровичу")
        for r in p.runs:
            set_run_ru(r)

# =========================================================================
# 7. REPLACE TABLE 3 (RASPISKA SIGNATURES) WITH FULL-WIDTH ROW FORMAT
# =========================================================================
tbl3 = doc.tables[2] # now table index 2
tbl3_elem = tbl3._tbl
parent3 = tbl3_elem.getparent()
idx3 = parent3.index(tbl3_elem)

elems_rasp = [
    make_p_xml("Деньги получил (Продавец):", size_pt=9.5, bold=True, space_before=2, space_after=1, keep_next=True),
    make_p_xml(line_str, size_pt=9.5, space_before=0, space_after=0, keep_next=True),
    make_p_xml("ФИО                                                                                                                                           подпись", size_pt=8.0, italic=True, space_before=0, space_after=8, keep_next=True),
    make_p_xml("Деньги передал (Покупатель):", size_pt=9.5, bold=True, space_before=2, space_after=1, keep_next=True),
    make_p_xml(line_str, size_pt=9.5, space_before=0, space_after=0, keep_next=True),
    make_p_xml("ФИО                                                                                                                                           подпись", size_pt=8.0, italic=True, space_before=0, space_after=4, keep_next=False),
]

for el in elems_rasp:
    parent3.insert(idx3, el)
    idx3 += 1

parent3.remove(tbl3_elem)

# Save updated DKP
out_dkp = os.path.join(base, "2026_10_23_ДКП_с_Распиской_Лазутина_Медведев.docx")
doc.save(out_dkp)
print("Saved DKP with CLEAN FULL-WIDTH ROW SIGNATURES:", out_dkp)
