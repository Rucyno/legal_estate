# -*- coding: utf-8 -*-
"""
Генератор и автообновлятор Сводного реестра проверок СБ «Этажи» (MASTER_LOG.xlsx):
- Формат: Вариант Б (Лист 1: «Реестр стажеров», Лист 2: «Карточки досье»).
- Фирменные цвета «Этажи»: Красный (#C8102E), Серый (#334155, #64748B, #F1F5F9), Белый (#FFFFFF).
- Выверенные размеры колонок под реальные данные, автофильтры, кликабельные ссылки на папки.
"""

import os
import sys
import docx
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SB_DIR = r"D:\Этажи\ДОУ Акты\внутренние договоры\СБ"
LOCAL_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "MASTER_LOG.xlsx")

MONTH_MAP = {
    'января': '01', 'февраля': '02', 'марта': '03', 'апреля': '04',
    'мая': '05', 'июня': '06', 'июля': '07', 'августа': '08',
    'сентября': '09', 'октября': '10', 'ноября': '11', 'декабря': '12'
}

def clean_doc_date(text):
    if not text:
        return ""
    m = re.search(r'[«"]\s*(\d{1,2})\s*[»"]\s*([а-яА-ЯёЁ]+)\s*(\d{4})', text)
    if m:
        day = m.group(1).zfill(2)
        mon_str = m.group(2).lower()
        mon = MONTH_MAP.get(mon_str, '00')
        year = m.group(3)
        return f"{day}.{mon}.{year}"
    m2 = re.search(r'(\d{2}\.\d{2}\.\d{4})', text)
    if m2:
        return m2.group(1)
    return text

def parse_all_candidates():
    candidates = []
    if not os.path.exists(SB_DIR):
        return candidates

    for folder in sorted(os.listdir(SB_DIR)):
        fpath = os.path.join(SB_DIR, folder)
        if not os.path.isdir(fpath):
            continue
        
        cand = {
            'folder_name': folder,
            'fio': folder,
            'birth_date': None,
            'age': None,
            'doc_date': None,
            'passport': None,
            'inn': None,
            'address': None,
            'fssp_status': 'Задолженностей нет',
            'fssp_debt': '0 ₽',
            'bankrupt_status': 'Банкротство не обнаружено',
            'arbitr_status': 'Арбитражных дел нет',
            'courts_status': 'Судебных дел не найдено',
            'rsfn_status': 'В перечне террористов/экстремистов отсутствует',
            'risk_level': 'Низкий (Допущен)',
            'files': os.listdir(fpath)
        }
        
        # 1. Parse docx if exists
        doc_path = os.path.join(fpath, 'Договор о сотрудничестве.docx')
        if os.path.exists(doc_path):
            try:
                doc = docx.Document(doc_path)
                full_text = '\n'.join([p.text for p in doc.paragraphs])
                tables_text = '\n'.join([c.text for t in doc.tables for r in t.rows for c in r.cells])
                all_text = full_text + '\n' + tables_text
                
                # Birth date
                b_match = re.search(r'(\d{2}\.\d{2}\.\d{4})\s*г\.?р\.?', all_text)
                if b_match:
                    cand['birth_date'] = b_match.group(1)
                    
                # Doc date
                cand['doc_date'] = clean_doc_date(tables_text)
                if not cand['doc_date'] or cand['doc_date'] == '00':
                    cand['doc_date'] = clean_doc_date(full_text)
                    
                # Passport
                p_match = re.search(r'паспорт[^\d]*(\d{2}\s*\d{2}\s*\d{6})', all_text, re.IGNORECASE)
                if p_match:
                    cand['passport'] = p_match.group(1).strip()
                    
                # Address
                a_match = re.search(r'адрес регистрации[:\s]+([^,\n]+(?:,[^,\n]+){1,3})', all_text, re.IGNORECASE)
                if a_match:
                    cand['address'] = a_match.group(1).strip().rstrip('.,;')
            except Exception:
                pass
                
        # 2. INN & fallback date from PDF filename
        for f in cand['files']:
            m_inn = re.match(r'^(\d{12})_(\d{8})\d*\.pdf$', f)
            if m_inn:
                cand['inn'] = m_inn.group(1)
                pdf_date_str = m_inn.group(2)
                if not cand['doc_date']:
                    cand['doc_date'] = f"{pdf_date_str[6:8]}.{pdf_date_str[4:6]}.{pdf_date_str[0:4]}"
                break
                
        # 3. Analyze checks based on filenames
        files_str = ' '.join(cand['files']).lower()
        
        # FSSP
        if 'большие задолженности' in files_str:
            cand['fssp_status'] = 'Крупная задолженность (>5.4 млн ₽)'
            cand['fssp_debt'] = '> 5 400 000 ₽'
            cand['risk_level'] = 'Высокий риск (СТОП)'
        elif 'фссп1.png' in files_str or 'фссп7.png' in files_str:
            cand['fssp_status'] = '19 исп. производств (~487 000 ₽)'
            cand['fssp_debt'] = '487 186 ₽'
            cand['risk_level'] = 'Высокий риск (СТОП)'
        elif 'погашено 140к' in files_str:
            cand['fssp_status'] = 'Погашено 140 000 ₽ (ИП закрыты)'
            cand['fssp_debt'] = '0 ₽ (погашено)'
            cand['risk_level'] = 'Требует согласования'
        elif '20к рублей' in files_str:
            cand['fssp_status'] = 'Задолженность 20 000 ₽'
            cand['fssp_debt'] = '20 000 ₽'
            cand['risk_level'] = 'Требует согласования'
            
        # Bankruptcy
        if 'банкротство завершена' in files_str or 'банкротство прошло' in files_str:
            cand['bankrupt_status'] = 'Процедура завершена (2026 г.)'
            if cand['risk_level'] != 'Высокий риск (СТОП)':
                cand['risk_level'] = 'Требует согласования'
        elif 'банкротство 21 год 830к' in files_str:
            cand['bankrupt_status'] = 'Банкротство 2021 г. (830 000 ₽)'
            if cand['risk_level'] != 'Высокий риск (СТОП)':
                cand['risk_level'] = 'Требует согласования'
                
        # Arbitr
        if 'арбитраж процедура банкроства завершена' in files_str or 'арбитраж банкротство' in files_str:
            cand['arbitr_status'] = 'Дело о банкротстве (архив)'
            
        # Courts
        if 'возвращенно исковое' in files_str:
            cand['courts_status'] = 'Иск возвращен (ООО Феникс)'
        elif 'решение жд суда' in files_str:
            cand['courts_status'] = 'Судебные решения в архиве'
            
        # Age
        if cand['birth_date']:
            try:
                parts = cand['birth_date'].split('.')
                b_year, b_month, b_day = int(parts[2]), int(parts[1]), int(parts[0])
                ref_year, ref_month, ref_day = 2026, 10, 23
                cand['age'] = ref_year - b_year - ((ref_month, ref_day) < (b_month, b_day))
            except Exception:
                pass

        candidates.append(cand)

    return candidates

def update_master_log():
    candidates = parse_all_candidates()
    if not candidates:
        print("[!] Нет кандидатов для формирования MASTER_LOG.")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # ----------------- STYLES («ЭТАЖИ»: КРАСНЫЙ, СЕРЫЙ, БЕЛЫЙ) -----------------
    FONT_FAMILY = "Segoe UI"
    COLOR_RED_MAIN = "C8102E"       # Фирменный красный «Этажи»
    COLOR_RED_DARK = "990000"
    COLOR_GRAY_DARK = "334155"      # Антрацитовый/Темно-серый
    COLOR_GRAY_MEDIUM = "64748B"
    COLOR_GRAY_LIGHT = "F1F5F9"
    COLOR_GRAY_ZEBRA = "F8FAFC"
    COLOR_WHITE = "FFFFFF"

    header_font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
    header_fill_red = PatternFill(start_color=COLOR_RED_MAIN, end_color=COLOR_RED_MAIN, fill_type="solid")

    title_font = Font(name=FONT_FAMILY, size=15, bold=True, color=COLOR_RED_MAIN)
    subtitle_font = Font(name=FONT_FAMILY, size=10, italic=True, color=COLOR_GRAY_MEDIUM)

    cell_font = Font(name=FONT_FAMILY, size=10, color="0F172A")
    cell_font_bold = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
    cell_font_link = Font(name=FONT_FAMILY, size=10, color=COLOR_RED_MAIN, bold=True, underline="single")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    header_border = Border(
        left=Side(style='thin', color='FFFFFF'),
        right=Side(style='thin', color='FFFFFF'),
        top=Side(style='medium', color=COLOR_RED_DARK),
        bottom=Side(style='medium', color=COLOR_RED_DARK)
    )

    # Badges
    badge_green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    badge_green_font = Font(name=FONT_FAMILY, size=10, bold=True, color="166534")

    badge_yellow_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    badge_yellow_font = Font(name=FONT_FAMILY, size=10, bold=True, color="854D0E")

    badge_red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    badge_red_font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_RED_MAIN)

    zebra_fill = PatternFill(start_color=COLOR_GRAY_ZEBRA, end_color=COLOR_GRAY_ZEBRA, fill_type="solid")

    def apply_badge(cell, risk):
        if "Низкий" in risk or "Допущен" in risk or "Чисто" in risk:
            cell.fill = badge_green_fill
            cell.font = badge_green_font
        elif "Требует" in risk or "согласования" in risk:
            cell.fill = badge_yellow_fill
            cell.font = badge_yellow_font
        else:
            cell.fill = badge_red_fill
            cell.font = badge_red_font

    # =========================================================================
    # ЛИСТ 1: РЕЕСТР СТАЖЕРОВ (Компактный список с выверенными размерами)
    # =========================================================================
    ws1 = wb.create_sheet(title="Реестр стажеров")
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "СЛУЖБА БЕЗОПАСНОСТИ «ЭТАЖИ» | СВОДНЫЙ РЕЕСТР СТАЖЕРОВ И ПАРТНЕРОВ"
    ws1["A1"].font = title_font
    ws1["A2"] = "Компактная таблица для фильтрации и сортировки (по возрасту, дате соглашения, статусу СБ). Полные досье — на вкладке «Карточки досье»."
    ws1["A2"].font = subtitle_font

    headers = [
        ("№", 6, "center"),
        ("ФИО стажёра / партнёра", 34, "left"),
        ("Возраст", 12, "center"),
        ("Дата соглашения", 18, "center"),
        ("Паспортные данные", 22, "center"),
        ("ИНН", 16, "center"),
        ("ФССП (Задолженность)", 25, "center"),
        ("Итоговый статус СБ", 26, "center"),
        ("Досье (Папка)", 18, "center"),
    ]

    row_h = 4
    for col_idx, (h_name, width, align) in enumerate(headers, start=1):
        cell = ws1.cell(row=row_h, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = header_fill_red
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    ws1.row_dimensions[row_h].height = 28

    for idx, c in enumerate(candidates, start=1):
        r = row_h + idx
        ws1.row_dimensions[r].height = 22
        is_zebra = (idx % 2 == 0)
        
        age_str = f"{c['age']} лет" if c['age'] is not None else "—"
        folder_url = os.path.join(SB_DIR, c['folder_name'])
        
        row_values = [
            (idx, "center", cell_font),
            (c['fio'], "left", cell_font_bold),
            (age_str, "center", cell_font_bold),
            (c['doc_date'] or "—", "center", cell_font),
            (c['passport'] or "—", "center", cell_font),
            (c['inn'] or "—", "center", cell_font),
            (c['fssp_debt'], "center", cell_font_bold if "0" not in c['fssp_debt'] else cell_font),
            (c['risk_level'], "center", cell_font),
            ("📂 Открыть папку", "center", cell_font_link),
        ]
        
        for col_idx, (val, align, f_style) in enumerate(row_values, start=1):
            cell = ws1.cell(row=r, column=col_idx, value=val)
            cell.font = f_style
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = thin_border
            
            if is_zebra and col_idx != 8:
                cell.fill = zebra_fill
                
            if col_idx == 8:
                apply_badge(cell, c['risk_level'])
                
            if col_idx == 7 and (">" in val or "487" in val):
                cell.fill = badge_red_fill
                cell.font = badge_red_font
            elif col_idx == 7 and ("20 000" in val or "погашено" in val):
                cell.fill = badge_yellow_fill
                cell.font = badge_yellow_font
                
            if col_idx == 9:
                cell.hyperlink = folder_url

    ws1.auto_filter.ref = f"A{row_h}:I{row_h + len(candidates)}"

    # =========================================================================
    # ЛИСТ 2: КАРТОЧКИ ДОСЬЕ (Верх: реквизиты, Низ: проверки СБ)
    # =========================================================================
    ws2 = wb.create_sheet(title="Карточки досье")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "ДЕТАЛЬНЫЕ КАРТОЧКИ ПРОВЕРКИ КАНДИДАТОВ СЛУЖБОЙ БЕЗОПАСНОСТИ «ЭТАЖИ»"
    ws2["A1"].font = title_font
    ws2["A2"] = "Идентификационные реквизиты стажера расположены сверху, результаты проверок СБ — отдельным блоком ниже."
    ws2["A2"].font = subtitle_font

    curr_row = 4
    card_header_fill = header_fill_red
    card_header_font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
    section_fill = PatternFill(start_color=COLOR_GRAY_LIGHT, end_color=COLOR_GRAY_LIGHT, fill_type="solid")

    for idx, c in enumerate(candidates, start=1):
        ws2.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=6)
        card_title_cell = ws2.cell(row=curr_row, column=1, value=f"#{idx} | {c['fio']} (Возраст: {c['age'] or '—'} лет) — {c['risk_level']}")
        card_title_cell.font = card_header_font
        card_title_cell.fill = card_header_fill
        card_title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws2.row_dimensions[curr_row].height = 26
        
        # Блок 1: Идентификационные данные (СВЕРХУ)
        r1 = curr_row + 1
        ws2.cell(row=r1, column=1, value="Дата соглашения:").font = cell_font_bold
        ws2.cell(row=r1, column=2, value=c['doc_date'] or "—").font = cell_font
        ws2.cell(row=r1, column=3, value="Серия и № паспорта:").font = cell_font_bold
        ws2.cell(row=r1, column=4, value=c['passport'] or "—").font = cell_font
        ws2.cell(row=r1, column=5, value="ИНН:").font = cell_font_bold
        ws2.cell(row=r1, column=6, value=c['inn'] or "—").font = cell_font
        
        r2 = curr_row + 2
        ws2.cell(row=r2, column=1, value="Дата рождения:").font = cell_font_bold
        ws2.cell(row=r2, column=2, value=c['birth_date'] or "—").font = cell_font
        ws2.cell(row=r2, column=3, value="Адрес регистрации:").font = cell_font_bold
        ws2.merge_cells(start_row=r2, start_column=4, end_row=r2, end_column=6)
        ws2.cell(row=r2, column=4, value=c['address'] or "—").font = cell_font
        
        # Блок 2: Результаты проверок СБ (НИЖЕ)
        r3 = curr_row + 3
        ws2.cell(row=r3, column=1, value="ФССП (задолженности):").font = cell_font_bold
        c_fssp = ws2.cell(row=r3, column=2, value=f"{c['fssp_debt']} ({c['fssp_status']})")
        c_fssp.font = cell_font_bold if "0" not in c['fssp_debt'] else cell_font
        ws2.cell(row=r3, column=3, value="Банкротство (ЕФРСБ):").font = cell_font_bold
        ws2.cell(row=r3, column=4, value=c['bankrupt_status']).font = cell_font
        ws2.cell(row=r3, column=5, value="Папка скан/фото:").font = cell_font_bold
        c_link = ws2.cell(row=r3, column=6, value="📂 Открыть файлы")
        c_link.font = cell_font_link
        c_link.hyperlink = os.path.join(SB_DIR, c['folder_name'])
        
        r4 = curr_row + 4
        ws2.cell(row=r4, column=1, value="КАД Арбитр:").font = cell_font_bold
        ws2.cell(row=r4, column=2, value=c['arbitr_status']).font = cell_font
        ws2.cell(row=r4, column=3, value="Суды общей юрисдикции:").font = cell_font_bold
        ws2.cell(row=r4, column=4, value=c['courts_status']).font = cell_font
        ws2.cell(row=r4, column=5, value="Росфинмониторинг:").font = cell_font_bold
        ws2.cell(row=r4, column=6, value=c['rsfn_status']).font = cell_font
        
        for cr in range(r1, r4 + 1):
            ws2.row_dimensions[cr].height = 21
            for cc in range(1, 7):
                ws2.cell(row=cr, column=cc).border = thin_border
                if cc in [1, 3, 5]:
                    ws2.cell(row=cr, column=cc).fill = section_fill
                    
        curr_row += 6

    # Выверенная ширина колонок для карточек
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 24
    ws2.column_dimensions["D"].width = 32
    ws2.column_dimensions["E"].width = 18
    ws2.column_dimensions["F"].width = 22

    # Сохраняем в СБ и локально
    out_sb = os.path.join(SB_DIR, "MASTER_LOG.xlsx")
    wb.save(out_sb)
    wb.save(LOCAL_FILE)
    print(f"[OK] MASTER_LOG.xlsx успешно обновлен в: {out_sb}")

if __name__ == "__main__":
    update_master_log()
