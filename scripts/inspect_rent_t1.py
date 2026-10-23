# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\10_Аренда_Греченкова_Артамонов\готовые"
f = os.path.join(base, "Договор_аренды_Греченкова_Артамонов_УСИЛЕННЫЙ_ПРОЕКТ.docx")
doc = docx.Document(f)

t1 = doc.tables[1]
print("--- CELL 0 (Арендодатель) ---")
for p in t1.cell(0, 0).paragraphs:
    print(repr(p.text))

print("--- CELL 1 (Арендатор) ---")
for p in t1.cell(0, 1).paragraphs:
    print(repr(p.text))
