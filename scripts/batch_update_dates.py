# -*- coding: utf-8 -*-
import docx, os, shutil

folder = r"C:\Users\PC\Desktop\Тёма"
archive_dir = os.path.join(folder, "archive")
os.makedirs(archive_dir, exist_ok=True)

# 1. Update Агентский_договор_Покупка.docx
doc_agent = docx.Document(os.path.join(folder, "Агентский_договор_Покупка.docx"))
# Table 1 [0,1]
cell = doc_agent.tables[1].rows[0].cells[1]
for p in cell.paragraphs:
    for r in p.runs:
        if '"13" октября 2026 г.' in r.text:
            r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')
doc_agent.save(os.path.join(folder, "Агентский_договор_Покупка.docx"))
print("Updated Агентский_договор_Покупка.docx")

# 2. Update Доп соглашение
doc_ds = docx.Document(os.path.join(folder, "2026_10_13_Горина_Ивлиев_Доп_соглашение_по_покупке.docx"))
# Table 0 [0,1]
for p in doc_ds.tables[0].rows[0].cells[1].paragraphs:
    for r in p.runs:
        if '"13" октября 2026 г.' in r.text:
            r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')
# Table 1 [0,1]
for p in doc_ds.tables[1].rows[0].cells[1].paragraphs:
    for r in p.runs:
        if '"13" октября 2026 г.' in r.text:
            r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')
# P7
for r in doc_ds.paragraphs[7].runs:
    if '"13" октября 2026 г.' in r.text:
        r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')

new_ds_path = os.path.join(folder, "2026_10_14_Горина_Ивлиев_Доп_соглашение_по_покупке.docx")
doc_ds.save(new_ds_path)
print("Updated & saved", new_ds_path)

# 3. Update Предварительный договор
doc_pd = docx.Document(os.path.join(folder, "2026_10_13_Горина_Ивлиев_Предварительный_договор.docx"))
# Table 0 [0,1]
for p in doc_pd.tables[0].rows[0].cells[1].paragraphs:
    for r in p.runs:
        if '"13" октября 2026 г.' in r.text:
            r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')
# P9 (1.3)
for r in doc_pd.paragraphs[9].runs:
    if '"13" ноября 2026 г.' in r.text:
        r.text = r.text.replace('"13" ноября 2026 г.', '"14" ноября 2026 г.')

new_pd_path = os.path.join(folder, "2026_10_14_Горина_Ивлиев_Предварительный_договор.docx")
doc_pd.save(new_pd_path)
print("Updated & saved", new_pd_path)

# 4. Update Расписка
doc_rec = docx.Document(os.path.join(folder, "2026_10_13_Горина_Ивлиев_Расписка.docx"))
# Table 0 [0,1]
for p in doc_rec.tables[0].rows[0].cells[1].paragraphs:
    for r in p.runs:
        if '"13" октября 2026 г.' in r.text:
            r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')
# P6
for r in doc_rec.paragraphs[6].runs:
    if '"13" октября 2026 г.' in r.text:
        r.text = r.text.replace('"13" октября 2026 г.', '"14" октября 2026 г.')

new_rec_path = os.path.join(folder, "2026_10_14_Горина_Ивлиев_Расписка.docx")
doc_rec.save(new_rec_path)
print("Updated & saved", new_rec_path)

# Move old 2026_10_13 files to archive
for f in ["2026_10_13_Горина_Ивлиев_Доп_соглашение_по_покупке.docx",
          "2026_10_13_Горина_Ивлиев_Предварительный_договор.docx",
          "2026_10_13_Горина_Ивлиев_Расписка.docx"]:
    src = os.path.join(folder, f)
    dst = os.path.join(archive_dir, f)
    if os.path.exists(src):
        shutil.move(src, dst)
        print(f"Moved {f} -> archive")

print("All updates completed successfully!")
