# -*- coding: utf-8 -*-
"""
Генератор соглашений о расторжении предварительных договоров купли-продажи
Разработано для LegalEstate под стандарты Ильи Сергеевича Фролова.
Строго 1 страница А4, Правило 10 и 11 (полноширинные границы подписей, без разрывов).
"""
import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

def set_run_ru(run, font_name="Times New Roman", size_pt=9.5, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    rPr = run._r.get_or_add_rPr()
    if not rPr.xpath("./w:lang"):
        rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:bidi="ru-RU"/>'))

def add_p(doc, text="", size_pt=9.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          space_after=2.0, line_spacing=1.06, indent=21.25, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Pt(indent)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if text:
        run = p.add_run(text)
        set_run_ru(run, size_pt=size_pt, bold=bold, italic=italic)
    return p

def add_header_table(doc, city_text="г. Пенза", date_text="«28» октября 2026 г."):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for cell in tbl.rows[0].cells:
        cell.width = Pt(260)
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r"""
            <w:tcBorders %s >
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        """ % nsdecls("w"))
        tcPr.append(tcBorders)

    p_l = tbl.cell(0, 0).paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.space_after = Pt(4.0)
    r_l = p_l.add_run(city_text)
    set_run_ru(r_l, size_pt=10.0, bold=True)

    p_r = tbl.cell(0, 1).paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.space_after = Pt(4.0)
    r_r = p_r.add_run(date_text)
    set_run_ru(r_r, size_pt=10.0, bold=True)

def add_sig_block(doc, title_text, signer_subtext="", keep_all=True):
    # Paragraph 1: Title
    add_p(doc, title_text, size_pt=10.0, bold=True, indent=0, space_after=2.0, keep_next=True)
    
    # Paragraph 2: Full width line (paragraph with bottom border)
    p_line = doc.add_paragraph()
    pPr = p_line._p.get_or_add_pPr()
    pBdr = parse_xml(f"""
    <w:pBdr {nsdecls("w")}>
        <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>
    </w:pBdr>
    """)
    pPr.append(pBdr)
    p_line.paragraph_format.space_after = Pt(1.5)
    p_line.paragraph_format.keep_with_next = True
    r_empty = p_line.add_run(" ")
    set_run_ru(r_empty, size_pt=9.5)
    
    # Paragraph 3: Subtext with right tab
    p_sub = doc.add_paragraph()
    pPr_sub = p_sub._p.get_or_add_pPr()
    tabs = parse_xml(f"""
    <w:tabs {nsdecls("w")}>
        <w:tab w:val="right" w:pos="9640"/>
    </w:tabs>
    """)
    pPr_sub.append(tabs)
    p_sub.paragraph_format.space_after = Pt(2.0 if signer_subtext else (6.0 if keep_all else 2.0))
    p_sub.paragraph_format.keep_with_next = True
    
    r1 = p_sub.add_run("ФИО полностью")
    set_run_ru(r1, size_pt=8.0, italic=True)
    r2 = p_sub.add_run("\tподпись")
    set_run_ru(r2, size_pt=8.0, italic=True)

    if signer_subtext:
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(6.0 if keep_all else 2.0)
        if keep_all:
            p_desc.paragraph_format.keep_with_next = True
        r_desc = p_desc.add_run(signer_subtext)
        set_run_ru(r_desc, size_pt=8.5, italic=True)

def generate_stroiteli_doc(out_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Pt(30.0)
        s.bottom_margin = Pt(30.0)
        s.left_margin = Pt(40.0)
        s.right_margin = Pt(40.0)

    # Title
    add_p(doc, "Соглашение о расторжении", size_pt=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=1.0)
    add_p(doc, "Предварительного договора купли-продажи квартиры от «20» октября 2026 г.", size_pt=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=4.0)

    add_header_table(doc, "г. Пенза", "«28» октября 2026 г.")

    add_p(doc, "Мы, нижеподписавшиеся:", size_pt=9.5, bold=True, indent=21.25, space_after=1.5)

    # Seller preambula
    seller_text = (
        "гр. Шатова Наталья Николаевна, 03.01.1956 года рождения, пол: женский, паспорт гражданина РФ: "
        "серия 5602 № 580464, выдан 16.03.2002 г. Отделом внутренних дел Октябрьского района г. Пензы, "
        "код подразделения: 582-002, зарегистрированная по адресу: Пензенская область, г. Пенза, пр-кт Строителей, д. 8, кв. 69, "
        "действующая от своего имени (как собственник 1/2 доли в праве), а также от имени и в интересах: "
        "гр. Валиевой Дарьи Геннадьевны, 20.06.1995 года рождения (собственника 1/4 доли в праве, паспорт РФ: серия 4519 № 291502, выдан 09.07.2019 г. ГУ МВД России по г. Москве, код подразделения: 770-038, адрес регистрации: Московская обл., г. Люберцы, ул. Весенняя, д. 12, кв. 142) "
        "на основании доверенности от 28.03.2026 г., серия 77АЕ № 2354924, удостоверенной нотариусом г. Москвы Авдеевой Е. И., зарегистрированной в реестре за № 77/1933-н/77-2026-3-392; "
        "гр. Смирновой Любови Геннадьевны, 08.05.1984 года рождения (собственника 1/4 доли в праве, паспорт РФ: серия 4015 № 353535, выдан 08.10.2015 г. ТП № 47 Отдела УФМС России по Санкт-Петербургу и Ленинградской обл. в Московском р-не г. Санкт-Петербурга, код подразделения: 780-047, адрес проживания: Пензенская обл., г. Пенза, пр-кт Строителей, д. 8, кв. 69) "
        "на основании доверенности от 21.08.2025 г., серия 58АА № 2314355, удостоверенной нотариусом г. Пензы Кочегиной В. В., зарегистрированной в реестре за № 58/26-н/58-2025-3-551, "
        "именуемые в дальнейшем совместно «Продавец», с одной стороны,"
    )
    add_p(doc, seller_text, size_pt=9.0, indent=21.25, space_after=2.0)

    # Buyer preambula
    buyer_text = (
        "и гр. Балакина Марина Олеговна, 21.07.1999 года рождения, пол: женский, паспорт гражданина РФ: "
        "серия 5621 № 642132, выдан 02.10.2021 г. УМВД России по Пензенской области, код подразделения: 580-001, "
        "зарегистрированная по адресу: Пензенская область, г. Пенза, ул. Красноармейская, д. 101, кв. 2, "
        "именуемая в дальнейшем «Покупатель», с другой стороны,"
    )
    add_p(doc, buyer_text, size_pt=9.0, indent=21.25, space_after=2.0)

    add_p(doc, "при совместном упоминании именуемые «Стороны», руководствуясь статьями 381, 421, 450, 451, 453 Гражданского кодекса РФ, заключили настоящее Соглашение о нижеследующем:", size_pt=9.0, indent=21.25, space_after=3.0)

    # Point 1
    p1 = (
        "1. В соответствии со статьями 450, 451 Гражданского кодекса РФ «Стороны» пришли к обоюдному соглашению "
        "расторгнуть Предварительный договор купли-продажи квартиры от «20» октября 2026 года (далее — «Договор»), "
        "предметом которого являлось обязательство заключить в будущем договор купли-продажи следующего объекта недвижимости: "
        "квартиры, назначение: жилое, площадь: 62,9 кв. м., этаж: 09, адрес: Пензенская область, г. Пенза, пр-кт Строителей, д. 8, кв. 69, "
        "кадастровый номер: 58:29:1007006:2498 (далее — «Объект»). С момента подписания настоящего Соглашения взаимные обязательства "
        "«Сторон», вытекающие из Договора, прекращаются в полном объеме."
    )
    add_p(doc, p1, size_pt=9.0, indent=21.25, space_after=2.5)

    # Point 2
    p2 = (
        "2. На основании пункта 11.2 Предварительного договора купли-продажи от «20» октября 2026 года и статей 381, 421 Гражданского кодекса РФ, "
        "в связи с неисполнением обязательств по инициативе «Покупателя» (отказом Покупателя от заключения основного договора купли-продажи), "
        "переданный «Покупателем» задаток в размере 50 000 (Пятьдесят тысяч) рублей 00 копеек (подтвержденный Распиской от «20» октября 2026 г.) "
        "остается в полном объеме у «Продавца» в качестве компенсации. «Покупатель» заявляет и подтверждает, что не имеет к «Продавцу» "
        "каких-либо имущественных или денежных претензий, связанных с невозвратом указанной суммы задатка."
    )
    add_p(doc, p2, size_pt=9.0, indent=21.25, space_after=2.5)

    # Point 3
    p3 = (
        "3. «Стороны» заявляют и подтверждают, что не имеют друг к другу каких-либо иных финансовых, материальных или имущественных "
        "претензий и требований (включая требования об уплате неустоек, возмещении убытков либо компенсации расходов), "
        "связанных с заключением, исполнением или расторжением Договора."
    )
    add_p(doc, p3, size_pt=9.0, indent=21.25, space_after=2.5)

    # Point 4
    p4 = (
        "4. Настоящее Соглашение составлено и подписано «Сторонами» в двух имеющих одинаковую юридическую силу экземплярах, "
        "по одному для каждой из «Сторон», вступает в силу с момента его подписания и является неотъемлемой частью Договора."
    )
    add_p(doc, p4, size_pt=9.0, indent=21.25, space_after=4.0, keep_next=True)

    # Signatures
    add_p(doc, "Подписи сторон:", size_pt=10.0, bold=True, indent=0, space_after=2.0, keep_next=True)
    
    seller_sub = "Шатова Наталья Николаевна, действующая от своего имени, а также от имени и в интересах Валиевой Дарьи Геннадьевны и Смирновой Любови Геннадьевны на основании доверенностей"
    add_sig_block(doc, "Продавец:", signer_subtext=seller_sub, keep_all=True)
    add_sig_block(doc, "Покупатель:", signer_subtext="Балакина Марина Олеговна", keep_all=False)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print("Document successfully created:", out_path)

def generate_central_doc(out_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Pt(32.0)
        s.bottom_margin = Pt(32.0)
        s.left_margin = Pt(40.0)
        s.right_margin = Pt(40.0)

    # Title
    add_p(doc, "Соглашение о расторжении", size_pt=12.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=1.0)
    add_p(doc, "Предварительного договора купли-продажи квартиры от «22» октября 2026 г.", size_pt=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=4.0)

    add_header_table(doc, "г. Пенза", "«27» октября 2026 г.")

    add_p(doc, "Мы, нижеподписавшиеся:", size_pt=9.5, bold=True, indent=21.25, space_after=1.5)

    seller_text = (
        "гр. Тюлюкина Любовь Даниловна, 19.02.1954 года рождения, пол: женский, паспорт гражданина РФ: "
        "серия 56 02 № 987708, выдан 06.02.2003 г. Отделом внутренних дел Первомайского района г. Пензы, "
        "код подразделения: 582-003, адрес регистрации: Пензенская область, г. Пенза, ул. Бригадная, д. 22, "
        "именуемая в дальнейшем «Продавец», с одной стороны,"
    )
    add_p(doc, seller_text, size_pt=9.5, indent=21.25, space_after=2.0)

    buyer_text = (
        "и гр. Лазутина Елена Викторовна, 10.09.1984 года рождения, пол: женский, паспорт гражданина РФ: "
        "серия 56 05 № 562191, выдан 05.05.2006 г. Отделом внутренних дел Первомайского района г. Пензы, "
        "код подразделения: 582-003, адрес регистрации: Пензенская область, Пензенский район, с. Засечное, ул. Олимпийская, д. 3, кв. 129, "
        "именуемая в дальнейшем «Покупатель», с другой стороны,"
    )
    add_p(doc, buyer_text, size_pt=9.5, indent=21.25, space_after=2.0)

    add_p(doc, "при совместном упоминании именуемые «Стороны», руководствуясь статьями 381, 421, 450, 451, 453 Гражданского кодекса РФ, заключили настоящее Соглашение о нижеследующем:", size_pt=9.5, indent=21.25, space_after=3.0)

    p1 = (
        "1. В соответствии со статьями 450, 451 Гражданского кодекса РФ «Стороны» пришли к обоюдному соглашению "
        "расторгнуть Предварительный договор купли-продажи квартиры от «22» октября 2026 года (далее — «Договор»), "
        "предметом которого являлась купля-продажа 1-комнатной квартиры, общей площадью: 44,8 кв. м., расположенной на 1 этаже "
        "по адресу: Российская Федерация, Пензенская область, Пензенский район, с.п. Засечный сельсовет, с. Засечное, ул. Центральная, д. 3, кв. 30, "
        "кадастровый номер: 58:24:0381402:206 (далее — «Объект»). С момента подписания настоящего Соглашения взаимные обязательства «Сторон», "
        "вытекающие из Договора, прекращаются в полном объеме."
    )
    add_p(doc, p1, size_pt=9.5, indent=21.25, space_after=2.5)

    p2 = (
        "2. На основании статей 381, 421 Гражданского кодекса РФ и условий Предварительного договора от «22» октября 2026 года, "
        "в связи с расторжением Договора по инициативе «Покупателя» (отказом Покупателя от заключения основного договора купли-продажи), "
        "переданный «Покупателем» задаток в размере 50 000 (Пятьдесят тысяч) рублей 00 копеек (подтвержденный распиской от «22» октября 2026 г.) "
        "остается в полном объеме у «Продавца» в качестве компенсации. «Покупатель» заявляет и подтверждает, что не имеет к «Продавцу» "
        "каких-либо имущественных или денежных претензий, связанных с невозвратом указанной суммы задатка."
    )
    add_p(doc, p2, size_pt=9.5, indent=21.25, space_after=2.5)

    p3 = (
        "3. «Стороны» заявляют и подтверждают, что не имеют друг к другу каких-либо иных финансовых, материальных или имущественных "
        "претензий и требований (включая требования об уплате неустоек, возмещении убытков либо компенсации расходов), "
        "связанных с заключением, исполнением или расторжением Договора."
    )
    add_p(doc, p3, size_pt=9.5, indent=21.25, space_after=2.5)

    p4 = (
        "4. Настоящее Соглашение составлено и подписано «Сторонами» в двух имеющих одинаковую юридическую силу экземплярах, "
        "по одному для каждой из «Сторон», вступает в силу с момента его подписания и является неотъемлемой частью Договора."
    )
    add_p(doc, p4, size_pt=9.5, indent=21.25, space_after=4.0, keep_next=True)

    add_p(doc, "Подписи сторон:", size_pt=10.0, bold=True, indent=0, space_after=2.0, keep_next=True)
    add_sig_block(doc, "Продавец:", signer_subtext="Тюлюкина Любовь Даниловна", keep_all=True)
    add_sig_block(doc, "Покупатель:", signer_subtext="Лазутина Елена Викторовна", keep_all=False)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print("Central document successfully created:", out_path)

if __name__ == "__main__":
    out1 = r"D:\Антигравити\legal_estate\deals\13_Шатова_Балакина_Строителей_8\2026_10_28_Соглашение_о_расторжении_Строителей_8_Шатова_Балакина.docx"
    out2 = r"D:\Антигравити\legal_estate\deals\13_Шатова_Балакина_Строителей_8\2026_10_27_Соглашение_о_расторжении_Центральная_3_Тюлюкина_Лазутина.docx"
    generate_stroiteli_doc(out1)
    generate_central_doc(out2)
