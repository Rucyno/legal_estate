# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
src_file = os.path.join(base, "Основной_договор_с_распиской_Лазутина_Умеркаева.docx")
doc = docx.Document(src_file)

print("Raspiska paragraphs:")
for i in range(48, len(doc.paragraphs)):
    print(f"P{i}: {repr(doc.paragraphs[i].text)}")

t3 = doc.tables[3]
print("Table 3 Cell 0:")
for idx, p in enumerate(t3.cell(0, 0).paragraphs):
    print(f"  {idx}: {repr(p.text)}")
print("Table 3 Cell 1:")
for idx, p in enumerate(t3.cell(0, 1).paragraphs):
    print(f"  {idx}: {repr(p.text)}")
