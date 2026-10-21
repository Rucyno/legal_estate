import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, indent=0.35, before=0, after=2, line_spacing=1.08, size=10):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size)
    return p

def create_receipt_doc(output_path):
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.59)
        s.right_margin = Inches(0.59)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.first_line_indent = Inches(0)
    p_t.paragraph_format.space_before = Pt(4)
    p_t.paragraph_format.space_after = Pt(2)
    r = p_t.add_run("РАСПИСКА")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.first_line_indent = Inches(0)
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("в получении полной оплаты стоимости квартиры по Договору купли-продажи квартиры от «22» октября 2026 года")
    r_sub.bold = True
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10)

    # Таблица города и даты
    table_head = doc.add_table(rows=1, cols=2)
    table_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_head.autofit = False
    tblPr = table_head._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)
    
    c1, c2 = table_head.cell(0, 0), table_head.cell(0, 1)
    c1.width = Inches(3.5); c2.width = Inches(3.5)
    set_cell_margins(c1, top=0, bottom=0, left=0, right=0)
    set_cell_margins(c2, top=0, bottom=0, left=0, right=0)
    
    p_city = c1.paragraphs[0]
    p_city.paragraph_format.space_before = Pt(0)
    p_city.paragraph_format.space_after = Pt(4)
    r_c = p_city.add_run("город Пенза")
    r_c.font.name = 'Times New Roman'; r_c.font.size = Pt(10); r_c.bold = True
    
    p_date = c2.paragraphs[0]
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(4)
    r_d = p_date.add_run("«22» октября 2026 года")
    r_d.font.name = 'Times New Roman'; r_d.font.size = Pt(10); r_d.bold = True

    add_p(doc, "Мы, нижеподписавшиеся:", indent=0.35, before=2, after=2)
    
    # Продавец 1: Асташкина В.С. за себя и за ребенка
    p_p1 = add_p(doc, indent=0.35, before=0, after=2)
    r = p_p1.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run("Асташкина Валерия Сергеевна")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run(
        ", 27.09.1995 года рождения, гражданство: Российская Федерация, пол: женский, "
        "паспорт гражданина РФ: серия 5619 № 536581, выдан 10.12.2019 г. УМВД России по Пензенской области, "
        "код подразделения: 580-011, адрес регистрации: Пензенская обл., Пензенский р-н, с. Засечное, ул. Изумрудная, д. 5, кв. 74, "
        "действующая за себя лично, а также от имени и в интересах своего малолетнего сына — гр. "
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run("Асташкина Константина Юрьевича")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run(
        ", 30.07.2020 года рождения, свидетельство о рождении: серия II-ИЗ № 613195, выдано 04.08.2020 г. "
        "Территориальным отделом ЗАГС Октябрьского района г. Пензы Управления ЗАГС Пензенской области, "
        "адрес регистрации: Пензенская обл., Пензенский р-н, с. Засечное, ул. Изумрудная, д. 5, кв. 74, и"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Продавец 2: Асташкин Ю.В.
    p_p2 = add_p(doc, indent=0.35, before=0, after=2)
    r = p_p2.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p2.add_run("Асташкин Юрий Владимирович")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p2.add_run(
        ", 08.11.1993 года рождения, гражданство: Российская Федерация, пол: мужской, "
        "паспорт гражданина РФ: серия 5613 № 171141, выдан 10.12.2013 г. отделением УФМС России по Пензенской области в Пензенском районе, "
        "код подразделения: 580-030, адрес регистрации: Пензенская обл., Пензенский р-н, с. Старая Каменка, ул. Быстровка, д. 21, "
        "именуемые совместно «Продавцы», настоящей распиской подтверждаем, что получили от:"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Покупатель: Петрова Н.Е.
    p_buy = add_p(doc, indent=0.35, before=0, after=2)
    r = p_buy.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_buy.add_run("Петровой Нины Евгеньевны")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_buy.add_run(
        ", 16.08.1965 года рождения, гражданство: Российская Федерация, пол: женский, "
        "паспорт гражданина РФ: серия 5610 № 948406, выдан 25.09.2010 г. отделением УФМС России по Пензенской области в Каменском районе, "
        "код подразделения: 580-005, адрес места пребывания: Пензенская обл., г. Пенза, ул. Кирова/Карла Маркса, д. 2а/19, кв. 13, "
        "именуемой «Покупатель»,"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Сумма и подтверждение
    add_p(doc, "денежные средства в размере 5 450 000 (Пять миллионов четыреста пятьдесят тысяч) рублей 00 копеек в качестве полной оплаты стоимости отчуждаемого жилого помещения — квартиры, расположенной по адресу: Пензенская область, г. Пенза, ул. Октябрьская, д. 4, кв. 17, общей площадью 65,9 кв.м., кадастровый номер: 58:29:2011001:823, в соответствии с условиями Договора купли-продажи квартиры от «22» октября 2026 года.", indent=0.35, before=2, after=2)

    add_p(doc, "Оплата произведена Покупателем наличными денежными средствами в полном объеме, из которых:", indent=0.35, before=0, after=2)
    add_p(doc, "- 50 000 (Пятьдесят тысяч) рублей 00 копеек выплачены ранее при подписании Предварительного договора купли-продажи от «24» августа 2026 года в качестве задатка;", indent=0.5, before=0, after=2)
    add_p(doc, "- 5 400 000 (Пять миллионов четыреста тысяч) рублей 00 копеек выплачены при подписании Договора купли-продажи квартиры от «22» октября 2026 года в качестве окончательного расчета.", indent=0.5, before=0, after=2)

    add_p(doc, "Обязательства Петровой Нины Евгеньевны по оплате стоимости вышеуказанной квартиры исполнены в полном объеме, своевременно и надлежащим образом. Расчет произведен полностью. Никаких финансовых, материальных или иных претензий имущественного характера по оплате отчуждаемой квартиры Продавцы к Петровой Нине Евгеньевне не имеют.", indent=0.35, before=2, after=4)

    add_p(doc, "Настоящая Расписка составлена в простой письменной форме в двух подлинных экземплярах, имеющих равную юридическую силу: один экземпляр хранится у Продавцов, один экземпляр выдан на руки Покупателю Петровой Нине Евгеньевне.", indent=0.35, before=0, after=4)

    # Таблица подписей
    t_sign = doc.add_table(rows=1, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.autofit = False
    tblPr = t_sign._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)
    c1, c2 = t_sign.cell(0, 0), t_sign.cell(0, 1)
    c1.width = Inches(3.6); c2.width = Inches(3.4)
    set_cell_margins(c1, top=40, bottom=40, left=0, right=20)
    set_cell_margins(c2, top=40, bottom=40, left=20, right=0)
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0); p1.paragraph_format.space_after = Pt(2); p1.paragraph_format.line_spacing = 1.05
    text_p1 = (
        "Денежные средства получили (Продавцы):\n\n"
        "____________________________________________\n"
        "Асташкина Валерия Сергеевна,\n"
        "действующая за себя лично и как законный представитель\n"
        "несовершеннолетнего Асташкина Константина Юрьевича\n\n"
        "____________________________________________\n"
        "Асташкин Юрий Владимирович\n"
    )
    r = p1.add_run(text_p1)
    r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2); p2.paragraph_format.line_spacing = 1.05
    text_p2 = (
        "Денежные средства передала (Покупатель):\n\n"
        "____________________________________________\n"
        "Петрова Нина Евгеньевна\n"
    )
    r = p2.add_run(text_p2)
    r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    doc.save(output_path)
    print("Receipt created successfully at:", output_path)


def create_transfer_act_doc(output_path):
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.59)
        s.bottom_margin = Inches(0.59)
        s.left_margin = Inches(0.59)
        s.right_margin = Inches(0.59)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.first_line_indent = Inches(0)
    p_t.paragraph_format.space_before = Pt(4)
    p_t.paragraph_format.space_after = Pt(2)
    r = p_t.add_run("ПЕРЕДАТОЧНЫЙ АКТ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.first_line_indent = Inches(0)
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("(АКТ ПРИЕМА-ПЕРЕДАЧИ КВАРТИРЫ)")
    r_sub.bold = True
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10)

    # Таблица города и даты
    table_head = doc.add_table(rows=1, cols=2)
    table_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_head.autofit = False
    tblPr = table_head._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)
    
    c1, c2 = table_head.cell(0, 0), table_head.cell(0, 1)
    c1.width = Inches(3.5); c2.width = Inches(3.5)
    set_cell_margins(c1, top=0, bottom=0, left=0, right=0)
    set_cell_margins(c2, top=0, bottom=0, left=0, right=0)
    
    p_city = c1.paragraphs[0]
    p_city.paragraph_format.space_before = Pt(0)
    p_city.paragraph_format.space_after = Pt(4)
    r_c = p_city.add_run("город Пенза")
    r_c.font.name = 'Times New Roman'; r_c.font.size = Pt(10); r_c.bold = True
    
    p_date = c2.paragraphs[0]
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(4)
    r_d = p_date.add_run("«22» октября 2026 года")
    r_d.font.name = 'Times New Roman'; r_d.font.size = Pt(10); r_d.bold = True

    add_p(doc, "Мы, нижеподписавшиеся:", indent=0.35, before=2, after=2)
    
    # Продавец 1
    p_p1 = add_p(doc, indent=0.35, before=0, after=2)
    r = p_p1.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run("Асташкина Валерия Сергеевна")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run(
        ", 27.09.1995 года рождения, гражданство: Российская Федерация, пол: женский, "
        "паспорт гражданина РФ: серия 5619 № 536581, выдан 10.12.2019 г. УМВД России по Пензенской области, "
        "код подразделения: 580-011, адрес регистрации: Пензенская обл., Пензенский р-н, с. Засечное, ул. Изумрудная, д. 5, кв. 74, "
        "действующая за себя лично, а также от имени и в интересах своего малолетнего сына — гр. "
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run("Асташкина Константина Юрьевича")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p1.add_run(
        ", 30.07.2020 года рождения, свидетельство о рождении: серия II-ИЗ № 613195, выдано 04.08.2020 г. "
        "Территориальным отделом ЗАГС Октябрьского района г. Пензы Управления ЗАГС Пензенской области, "
        "адрес регистрации: Пензенская обл., Пензенский р-н, с. Засечное, ул. Изумрудная, д. 5, кв. 74, и"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Продавец 2
    p_p2 = add_p(doc, indent=0.35, before=0, after=2)
    r = p_p2.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p2.add_run("Асташкин Юрий Владимирович")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_p2.add_run(
        ", 08.11.1993 года рождения, гражданство: Российская Федерация, пол: мужской, "
        "паспорт гражданина РФ: серия 5613 № 171141, выдан 10.12.2013 г. отделением УФМС России по Пензенской области в Пензенском районе, "
        "код подразделения: 580-030, адрес регистрации: Пензенская обл., Пензенский р-н, с. Старая Каменка, ул. Быстровка, д. 21, "
        "именуемые совместно «Продавец», с одной стороны, и"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Покупатель
    p_buy = add_p(doc, indent=0.35, before=0, after=2)
    r = p_buy.add_run("гр. ")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_buy.add_run("Петрова Нина Евгеньевна")
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    r = p_buy.add_run(
        ", 16.08.1965 года рождения, гражданство: Российская Федерация, пол: женский, "
        "паспорт гражданина РФ: серия 5610 № 948406, выдан 25.09.2010 г. отделением УФМС России по Пензенской области в Каменском районе, "
        "код подразделения: 580-005, адрес места пребывания: Пензенская обл., г. Пенза, ул. Кирова/Карла Маркса, д. 2а/19, кв. 13, "
        "именуемая в дальнейшем «Покупатель», с другой стороны, "
        "совместно именуемые «Стороны», составили настоящий Передаточный акт к Договору купли-продажи квартиры от «22» октября 2026 года о нижеследующем:"
    )
    r.font.name = 'Times New Roman'; r.font.size = Pt(10)

    # Пункты акта
    add_p(doc, "1. В соответствии со статьей 556 Гражданского кодекса Российской Федерации и условиями Договора купли-продажи квартиры от «22» октября 2026 года Продавец передал, а Покупатель принял в собственность жилое помещение — квартиру, расположенную по адресу: Пензенская область, г. Пенза, ул. Октябрьская, д. 4, кв. 17, общей площадью 65,9 (Шестьдесят пять целых девять десятых) кв.м., расположенную на 1 этаже жилого дома, кадастровый номер: 58:29:2011001:823 (далее по тексту — «Объект»).", indent=0.35, before=2, after=2)

    add_p(doc, "2. Одновременно с Объектом Продавец передал Покупателю следующее согласованное Сторонами имущество в исправном рабочем состоянии, стоимость которого полностью включена в общую цену Объекта недвижимости (п. 6.8 Предварительного договора):", indent=0.35, before=0, after=2)
    add_p(doc, "- шкаф;", indent=0.5, before=0, after=1)
    add_p(doc, "- комод;", indent=0.5, before=0, after=1)
    add_p(doc, "- полка настенная;", indent=0.5, before=0, after=1)
    add_p(doc, "- гардины и шторы;", indent=0.5, before=0, after=1)
    add_p(doc, "- кухонный гарнитур с газовой варочной поверхностью (без духового шкафа);", indent=0.5, before=0, after=1)
    add_p(doc, "- водонагреватель (бойлер).", indent=0.5, before=0, after=2)

    add_p(doc, "3. Объект осмотрен Покупателем до подписания настоящего Передаточного акта. Покупатель подтверждает, что фактическое техническое и санитарное состояние квартиры, конструктивных элементов, дверей, окон, напольных покрытий, системы отопления, электропроводки и санитарно-технического оборудования полностью соответствует условиям Договора купли-продажи и пригодно для постоянного проживания. Скрытых или явных недостатков, о которых не было бы заранее сообщено Продавцом, не обнаружено. Претензий по техническому состоянию и комплектации квартиры Покупатель к Продавцу не имеет.", indent=0.35, before=2, after=2)

    add_p(doc, "4. Одновременно с подписанием настоящего акта Продавец передал Покупателю полный комплект ключей от входной двери квартиры и подъезда в количестве: _____ шт., а также техническую документацию, паспорта на установленные приборы учета и квитанции, подтверждающие отсутствие задолженностей по коммунальным услугам, взносам на капитальный ремонт, электроснабжению, теплоснабжению, водоотведению и вывозу ТКО.", indent=0.35, before=0, after=2)

    add_p(doc, "5. Показания индивидуальных приборов учета (счетчиков) на момент фактической передачи квартиры зафиксированы Сторонами следующие:", indent=0.35, before=0, after=2)
    add_p(doc, "- Счетчик электроэнергии: № ________________________, показания: ________________ кВт*ч;", indent=0.5, before=0, after=1)
    add_p(doc, "- Счетчик холодного водоснабжения (ХВС): № ___________________, показания: ________________ куб.м;", indent=0.5, before=0, after=1)
    add_p(doc, "- Счетчик горячего водоснабжения (ГВС): № ___________________, показания: ________________ куб.м;", indent=0.5, before=0, after=1)
    add_p(doc, "- Счетчик газа: № ___________________, показания: ________________ куб.м.", indent=0.5, before=0, after=2)

    add_p(doc, "6. Взаимные финансовые расчеты между Сторонами произведены в полном объеме наличными денежными средствами в размере 5 450 000 (Пять миллионов четыреста пятьдесят тысяч) рублей 00 копеек в соответствии с Распиской от «22» октября 2026 года. Материальных и иных имущественных претензий Стороны друг к другу не имеют.", indent=0.35, before=2, after=2)

    add_p(doc, "7. С момента подписания настоящего Передаточного акта обязанность Продавца по фактической передаче квартиры и имущества Покупателю, а также обязанность Покупателя по принятию квартиры считаются исполненными полностью и надлежащим образом (п. 1 ст. 556 ГК РФ). Риск случайной гибели или повреждения квартиры переходит на Покупателя с момента подписания настоящего акта.", indent=0.35, before=0, after=2)

    add_p(doc, "8. Настоящий акт составлен и подписан в двух подлинных экземплярах, имеющих одинаковую юридическую силу: по одному экземпляру для каждой из Сторон.", indent=0.35, before=0, after=4)

    # Таблица подписей акта
    t_sign = doc.add_table(rows=1, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.autofit = False
    tblPr = t_sign._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)
    c1, c2 = t_sign.cell(0, 0), t_sign.cell(0, 1)
    c1.width = Inches(3.6); c2.width = Inches(3.4)
    set_cell_margins(c1, top=40, bottom=40, left=0, right=20)
    set_cell_margins(c2, top=40, bottom=40, left=20, right=0)
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0); p1.paragraph_format.space_after = Pt(2); p1.paragraph_format.line_spacing = 1.05
    text_p1 = (
        "СДАЛ (ПРОДАВЕЦ):\n\n"
        "____________________________________________\n"
        "Асташкина Валерия Сергеевна,\n"
        "действующая за себя лично и как законный представитель\n"
        "несовершеннолетнего Асташкина Константина Юрьевича\n\n"
        "____________________________________________\n"
        "Асташкин Юрий Владимирович\n\n"
        "Дата: «22» октября 2026 года\n"
    )
    r = p1.add_run(text_p1)
    r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2); p2.paragraph_format.line_spacing = 1.05
    text_p2 = (
        "ПРИНЯЛ (ПОКУПАТЕЛЬ):\n\n"
        "____________________________________________\n"
        "Петрова Нина Евгеньевна\n\n\n\n\n"
        "Дата: «22» октября 2026 года\n"
    )
    r = p2.add_run(text_p2)
    r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)

    doc.save(output_path)
    print("Transfer Act created successfully at:", output_path)

if __name__ == '__main__':
    out_dir = 'deals/11_Асташкины_Петрова/готовые'
    os.makedirs(out_dir, exist_ok=True)
    
    f_receipt = os.path.join(out_dir, '2026_10_22_Расписка_полная_оплата_Асташкины_Петрова.docx')
    f_act = os.path.join(out_dir, '2026_10_22_Передаточный_акт_Асташкины_Петрова.docx')
    
    create_receipt_doc(f_receipt)
    create_transfer_act_doc(f_act)
    print("ALL DEAL 11 DOCUMENTS CREATED SUCCESSFULLY!")
