# -*- coding: utf-8 -*-
import docx, os, sys
sys.stdout.reconfigure(encoding="utf-8")

base = r"D:\Антигравити\legal_estate\deals\07_Колесова\готовые"
f_dkp = os.path.join(base, "2026_10_23_ДКП_с_Распиской_Лазутина_Медведев.docx")
doc = docx.Document(f_dkp)

for i in range(35, 45):
    if i < len(doc.paragraphs):
        print(f"P{i}: {doc.paragraphs[i].text}")
