# -*- coding: utf-8 -*-
import docx, os, sys
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

# Let us test creating a signature block with bottom border or full line
doc = docx.Document()
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pBdr = parse_xml(f"""
<w:pBdr {nsdecls("w")}>
    <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>
</w:pBdr>
""")
pPr.append(pBdr)
p.add_run("Подпись: ")
doc.save("test_border.docx")
print("Saved test_border.docx")
